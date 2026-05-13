from __future__ import annotations

import argparse
import sys
from pathlib import Path

from openpyxl import load_workbook

try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass


def nz(value) -> str:
    return "" if value is None else str(value).strip()


def table_rows(ws) -> list[list[str]]:
    rows: list[list[str]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        values = [nz(cell) for cell in row]
        if any(values):
            rows.append(values)
    return rows


def build_route_summary(route_workbook: Path, factory_workbook: Path) -> list[dict]:
    route_wb = load_workbook(route_workbook, data_only=True)
    factory_wb = load_workbook(factory_workbook, data_only=True)
    try:
        material_rows = table_rows(route_wb.worksheets[1])
        route_rows = table_rows(route_wb.worksheets[5])
        op_rows = table_rows(route_wb.worksheets[6])
        mat_link_rows = table_rows(route_wb.worksheets[8])
        step_rows = table_rows(route_wb.worksheets[9])
        wc_rows = table_rows(factory_wb.worksheets[7])

        material_name = {row[24]: row[1] for row in material_rows if len(row) > 24 and row[24]}
        wc_name = {row[3]: row[0] for row in wc_rows if len(row) > 3 and row[3]}

        routes: list[dict] = []
        for route_row in route_rows:
            route_code = route_row[8]
            route_ops = [row for row in op_rows if len(row) > 13 and row[13] == route_code]
            route_ops.sort(key=lambda row: int(row[0]))
            ops: list[dict] = []
            for op_row in route_ops:
                op_no = op_row[0]
                op_mat_rows = [row for row in mat_link_rows if len(row) > 4 and row[1] == route_code and row[4] == op_no]
                op_mat_rows.sort(key=lambda row: row[3])
                op_step_rows = [row for row in step_rows if len(row) > 2 and row[1] == route_code and row[2] == op_no]
                op_step_rows.sort(key=lambda row: int(row[3]))
                ops.append(
                    {
                        "op_no": op_row[0],
                        "op_type": op_row[1],
                        "op_name": op_row[14],
                        "content": op_row[2],
                        "wc_code": op_row[3],
                        "wc_name": wc_name.get(op_row[3], ""),
                        "assist_time": op_row[7],
                        "run_time": op_row[8],
                        "output_code": op_row[6],
                        "output_name": material_name.get(op_row[6], ""),
                        "prev_op": op_row[15],
                        "materials": [
                            {
                                "code": row[3],
                                "name": material_name.get(row[3], ""),
                                "qty": row[5],
                            }
                            for row in op_mat_rows
                        ],
                        "steps": [
                            {
                                "no": row[3],
                                "name": row[4],
                                "content": row[5],
                            }
                            for row in op_step_rows
                        ],
                    }
                )

            total_run_time = sum(int(op["run_time"]) for op in ops if op["run_time"].isdigit())
            routes.append(
                {
                    "route_name": route_row[0],
                    "route_type": route_row[1],
                    "material_code": route_row[4],
                    "material_name": material_name.get(route_row[4], ""),
                    "version": route_row[7],
                    "route_code": route_row[8],
                    "factory_org": route_row[11],
                    "total_run_time": total_run_time,
                    "ops": ops,
                }
            )
        return routes
    finally:
        route_wb.close()
        factory_wb.close()


def route_filename(route_name: str) -> str:
    return f"技术文件-{route_name}.md"


def render_material_table(materials: list[dict]) -> str:
    if not materials:
        return "无\n"
    lines = [
        "| 物料名称 | 物料编码 | 数量 |",
        "| --- | --- | ---: |",
    ]
    for item in materials:
        lines.append(f"| {item['name'] or '-'} | `{item['code']}` | {item['qty']} |")
    return "\n".join(lines) + "\n"


def render_steps(steps: list[dict]) -> str:
    if not steps:
        return "- 无\n"
    lines = []
    for step in steps:
        lines.append(f"- `{step['name']}`: {step['content']}")
    return "\n".join(lines) + "\n"


def render_route_doc(route: dict, source_name: str, route_workbook_name: str) -> str:
    flow = " -> ".join(f"`{op['op_no']} {op['op_name']}`" for op in route["ops"])
    lines = [
        f"# {route['route_name']}技术文件",
        "",
        "## 基本信息",
        f"- 工艺路线名称: {route['route_name']}",
        f"- 工艺路线编码: `{route['route_code']}`",
        f"- 工艺路线版本: `{route['version']}`",
        f"- 工艺类型: `{route['route_type']}`",
        f"- 目标物料: {route['material_name']} (`{route['material_code']}`)",
        f"- 工厂组织: `{route['factory_org']}`",
        f"- 工序数量: `{len(route['ops'])}`",
        f"- 总定额加工工时: `{route['total_run_time']} 分钟`",
        "",
        "## 工艺流程",
        flow,
        "",
        "## 工序一览",
        "| 工序号 | 工序名称 | 工序类型 | 工作中心 | 定额加工工时(分钟) | 产出物料 |",
        "| --- | --- | --- | --- | ---: | --- |",
    ]
    for op in route["ops"]:
        lines.append(
            f"| {op['op_no']} | {op['op_name']} | {op['op_type']} | {op['wc_name']} (`{op['wc_code']}`) | {op['run_time']} | {op['output_name']} (`{op['output_code']}`) |"
        )

    lines.extend(["", "## 工序说明"])
    for op in route["ops"]:
        lines.extend(
            [
                "",
                f"### {op['op_no']} {op['op_name']}",
                f"- 工序类型: `{op['op_type']}`",
                f"- 工作中心: {op['wc_name']} (`{op['wc_code']}`)",
                f"- 前置工序: `{op['prev_op'] or '无'}`",
                f"- 定额加工工时: `{op['run_time']} 分钟`",
                f"- 定额辅助工时: `{op['assist_time']} 分钟`",
                f"- 产出物料: {op['output_name']} (`{op['output_code']}`)",
                f"- 工序内容: {op['content'] or '无'}",
                "",
                "工序物料:",
                render_material_table(op["materials"]).rstrip(),
                "",
                "关键控制与记录:",
                render_steps(op["steps"]).rstrip(),
            ]
        )

    lines.extend(
        [
            "",
            "## 生成依据",
            f"- 来源样例: `{source_name}`",
            f"- 工艺路线数据: `{route_workbook_name}`",
        ]
    )
    return "\n".join(lines) + "\n"


def parse_args():
    parser = argparse.ArgumentParser(description="根据最新工艺路线工作簿导出 Markdown 技术文件")
    parser.add_argument("--route-workbook", type=Path, required=True, help="产品与工艺工作簿")
    parser.add_argument("--factory-workbook", type=Path, required=True, help="工厂资源工作簿")
    parser.add_argument("--source-name", default="大疆无人机BOM与工艺路线.xlsx", help="来源样例名称")
    parser.add_argument("--output-dir", type=Path, default=None, help="输出目录，默认与工艺路线工作簿同目录")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    output_dir = (args.output_dir or args.route_workbook.parent).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    routes = build_route_summary(args.route_workbook.resolve(), args.factory_workbook.resolve())
    for route in routes:
        output_path = output_dir / route_filename(route["route_name"])
        output_path.write_text(
            render_route_doc(route, args.source_name, args.route_workbook.name),
            encoding="utf-8",
        )
        print(f"[OK] 已生成: {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
