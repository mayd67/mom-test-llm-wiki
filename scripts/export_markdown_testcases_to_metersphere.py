from __future__ import annotations

import argparse
import re
import shutil
from pathlib import Path

from openpyxl import load_workbook


HEADER = [
    "用例名称",
    "所属模块",
    "标签",
    "前置条件",
    "步骤描述",
    "预期结果",
    "编辑模式",
    "备注",
    "用例状态",
    "责任人",
    "用例等级",
]


def parse_table_row(line: str) -> list[str]:
    if not line.startswith("|") or not line.endswith("|"):
        raise ValueError(f"invalid table row: {line}")
    return [cell.strip() for cell in line.strip()[1:-1].split("|")]


def is_separator_row(cells: list[str]) -> bool:
    return all(cell.replace("-", "").replace(":", "").strip() == "" for cell in cells)


def append_cell_value(original: str, continuation: str) -> str:
    if not original:
        return continuation
    if not continuation:
        return original
    return f"{original}<br>{continuation}"


def extract_cases(markdown_path: Path) -> list[dict[str, str]]:
    cases: list[dict[str, str]] = []
    lines = markdown_path.read_text(encoding="utf-8").splitlines()

    for line in lines:
        if not line.startswith("|"):
            continue
        cells = parse_table_row(line)
        if len(cells) not in (11, 12):
            continue
        if cells[0] in ("用例编码", "用例名称") or is_separator_row(cells):
            continue

        if len(cells) == 12:
            code, name, module, tag, precondition, steps, expected, _mode, note, status, owner, level = cells
            case_name = f"{code} {name}".strip()
        else:
            case_name, module, tag, precondition, steps, expected, _mode, note, status, owner, level = cells

        # Markdown working drafts use blank shared columns for a case's later steps.
        if not case_name:
            if cases and (steps or expected):
                previous = cases[-1]
                previous["步骤描述"] = append_cell_value(previous["步骤描述"], steps)
                previous["预期结果"] = append_cell_value(previous["预期结果"], expected)
            continue

        cases.append(
            {
                "用例名称": case_name,
                "所属模块": module,
                "标签": tag,
                "前置条件": precondition,
                "步骤描述": steps,
                "预期结果": expected,
                # MeterSphere imports require TEXT mode, regardless of the Markdown work draft.
                "编辑模式": "TEXT",
                "备注": note,
                "用例状态": status or "未开始",
                "责任人": owner,
                "用例等级": level,
            }
        )
    return cases


def strip_case_code_prefix(name: str) -> str:
    return re.sub(r"^TC-[0-9A-Z]+[\s\u3000]+", "", name).strip()


def clear_template_rows(ws) -> None:
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)


def clear_merged_cells(ws) -> None:
    for merged_range in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merged_range))


def write_cases(ws, cases: list[dict[str, str]]) -> None:
    for row_idx, case in enumerate(cases, start=2):
        for col_idx, key in enumerate(HEADER, start=1):
            ws.cell(row=row_idx, column=col_idx, value=case[key] or None)


def validate_header(ws) -> None:
    current = [ws.cell(1, idx).value for idx in range(1, len(HEADER) + 1)]
    if current != HEADER:
        raise ValueError(f"unexpected template header: {current}")


def export(
    markdown_path: Path,
    template_path: Path,
    output_path: Path,
    *,
    strip_code_prefix: bool = False,
    clear_notes: bool = False,
    owner: str | None = None,
) -> int:
    cases = extract_cases(markdown_path)
    if not cases:
        raise ValueError(f"no cases parsed from {markdown_path}")

    for case in cases:
        if strip_code_prefix:
            case["用例名称"] = strip_case_code_prefix(case["用例名称"])
        if clear_notes:
            case["备注"] = ""
        if owner is not None:
            case["责任人"] = owner

    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(template_path, output_path)

    wb = load_workbook(output_path)
    ws = wb[wb.sheetnames[0]]
    validate_header(ws)
    clear_merged_cells(ws)
    clear_template_rows(ws)
    write_cases(ws, cases)
    wb.save(output_path)
    return len(cases)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Export markdown testcases to MeterSphere xlsx.")
    parser.add_argument("--input", required=True, type=Path, help="Markdown testcase file.")
    parser.add_argument("--template", required=True, type=Path, help="MeterSphere template xlsx.")
    parser.add_argument("--output", required=True, type=Path, help="Output xlsx path.")
    parser.add_argument("--strip-code-prefix", action="store_true", help="Strip TC code prefix from case name.")
    parser.add_argument("--clear-notes", action="store_true", help="Clear notes column.")
    parser.add_argument("--owner", default=None, help="Set owner for all rows.")
    return parser


def main() -> None:
    args = build_parser().parse_args()
    count = export(
        args.input,
        args.template,
        args.output,
        strip_code_prefix=args.strip_code_prefix,
        clear_notes=args.clear_notes,
        owner=args.owner,
    )
    print(f"exported {count} cases to {args.output}")


if __name__ == "__main__":
    main()
