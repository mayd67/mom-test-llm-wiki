from __future__ import annotations

import argparse
import importlib.util
import json
import os
import re
import sys
import zipfile
from collections import OrderedDict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass

ROOT = Path(__file__).resolve().parents[1]
TEST_DATA_OUTPUT_ROOT = ROOT / "outputs" / "06_测试数据"
SCENARIO_OUTPUT_ROOT = TEST_DATA_OUTPUT_ROOT / "大疆无人机"
DEFAULT_INPUT = ROOT / "raw" / "dji_uav_bom_route.xlsx"
DEFAULT_SEED = SCENARIO_OUTPUT_ROOT / "dji_uav_demo_seed.json"
DEFAULT_PACKAGE_DIR = SCENARIO_OUTPUT_ROOT / "dji_uav_demo_import"
DEFAULT_PACKAGE_ZIP = SCENARIO_OUTPUT_ROOT / "dji_uav_demo_import.zip"
DEFAULT_ROUTE_ONLY_WORKBOOK = SCENARIO_OUTPUT_ROOT / "dji_uav_route_only_import.xlsx"
VERSION = "A.01"
SECURITY = "内部"
RELEASE_TIME = "2026-05-12 08:00:00"
INTEGRATION_SYSTEM = "MOM-SEED"
ROUTE_ONLY_SHEETS = {"填写说明", "工艺路线", "工艺路线工序", "工艺路线工序序列", "工艺路线工序物料"}

ORG = {
    "admin_company": "ADM-DJI-HQ",
    "admin_tech": "ADM-DJI-TECH",
    "admin_pmc": "ADM-DJI-PMC",
    "admin_qa": "ADM-DJI-QA",
    "admin_wm": "ADM-DJI-WM",
    "admin_em": "ADM-DJI-EM",
    "admin_plant": "ADM-DJI-PLANT",
    "biz_company": "BIZ-DJI-COMPANY",
    "biz_tech": "BIZ-DJI-TECH",
    "biz_pmc": "BIZ-DJI-PMC",
    "biz_qa": "BIZ-DJI-QA",
    "biz_wm": "BIZ-DJI-WM",
    "biz_em": "BIZ-DJI-EM",
    "biz_plant": "BIZ-DJI-PLANT",
    "biz_asm": "BIZ-DJI-ASM-WS",
    "biz_smt": "BIZ-DJI-SMT-WS",
    "biz_test": "BIZ-DJI-TST-WS",
    "biz_pack": "BIZ-DJI-PKG-WS",
}

USERS = [
    ("100", "程亦凡", ORG["admin_tech"], ORG["biz_tech"], "工艺工程师"),
    ("101", "夏知遥", ORG["admin_pmc"], ORG["biz_pmc"], "计划员"),
    ("102", "林述安", ORG["admin_plant"], ORG["biz_smt"], "SMT工艺工程师"),
    ("103", "沈嘉树", ORG["admin_qa"], ORG["biz_qa"], "质量工程师"),
    ("104", "顾念慈", ORG["admin_wm"], ORG["biz_wm"], "仓库管理员"),
    ("105", "唐雨杭", ORG["admin_plant"], ORG["biz_asm"], "总装班组长"),
    ("106", "顾星野", ORG["admin_plant"], ORG["biz_test"], "调试测试工程师"),
    ("107", "许安歌", ORG["admin_plant"], ORG["biz_pack"], "包装操作工"),
]

WORK_CENTERS = OrderedDict(
    {
        "总装备料区": {
            "code": "WC-DJI-KIT",
            "org": ORG["biz_asm"],
            "type": "产线",
            "class": "加工",
        },
        "电装工位": {
            "code": "WC-DJI-ELEC",
            "org": ORG["biz_asm"],
            "type": "产线",
            "class": "加工",
        },
        "测试工位": {
            "code": "WC-DJI-TEST",
            "org": ORG["biz_test"],
            "type": "产线",
            "class": "检验",
        },
        "钳装工位": {
            "code": "WC-DJI-MECH",
            "org": ORG["biz_asm"],
            "type": "产线",
            "class": "加工",
        },
        "调试工位": {
            "code": "WC-DJI-DEBUG",
            "org": ORG["biz_test"],
            "type": "产线",
            "class": "检验",
        },
        "质量工位": {
            "code": "WC-DJI-QA",
            "org": ORG["biz_test"],
            "type": "产线",
            "class": "检验",
        },
        "包装工位": {
            "code": "WC-DJI-PACK",
            "org": ORG["biz_pack"],
            "type": "产线",
            "class": "加工",
        },
        "SMT备料区": {
            "code": "WC-DJI-SMT-KIT",
            "org": ORG["biz_smt"],
            "type": "产线",
            "class": "加工",
        },
        "SMT产线": {
            "code": "WC-DJI-SMT",
            "org": ORG["biz_smt"],
            "type": "产线",
            "class": "加工",
        },
        "清洗站": {
            "code": "WC-DJI-CLEAN",
            "org": ORG["biz_smt"],
            "type": "产线",
            "class": "加工",
        },
        "手工焊台区": {
            "code": "WC-DJI-PTH",
            "org": ORG["biz_smt"],
            "type": "产线",
            "class": "加工",
        },
        "PCB检验区": {
            "code": "WC-DJI-PCB-QA",
            "org": ORG["biz_smt"],
            "type": "产线",
            "class": "检验",
        },
    }
)

WC_USER_LINKS = {
    "WC-DJI-KIT": "105",
    "WC-DJI-ELEC": "105",
    "WC-DJI-TEST": "106",
    "WC-DJI-MECH": "105",
    "WC-DJI-DEBUG": "106",
    "WC-DJI-QA": "103",
    "WC-DJI-PACK": "107",
    "WC-DJI-SMT-KIT": "102",
    "WC-DJI-SMT": "102",
    "WC-DJI-CLEAN": "102",
    "WC-DJI-PTH": "102",
    "WC-DJI-PCB-QA": "103",
}

ROUTES = {
    "无人机总装工艺路线": {
        "code": "RT-DJI-UAV-ASM-A01",
        "name": "无人机总装工艺路线",
        "route_spec": "装配",
        "op_spec": "装配专业",
        "org": ORG["biz_asm"],
        "material_code": "DJI-UAV-01",
    },
    "电控印制板工艺路线": {
        "code": "RT-DJI-PCB-A01",
        "name": "电控印制板工艺路线",
        "route_spec": "装配",
        "op_spec": "装配专业",
        "org": ORG["biz_smt"],
        "material_code": "S-PCB-05",
    },
}

ROUTE_OUTPUTS = {
    "RT-DJI-UAV-ASM-A01": {},
    "RT-DJI-PCB-A01": {
        "20": "E-SMT-01",
        "30": "E-SMT-01",
        "40": "S-PCB-05",
        "50": "S-PCB-05",
        "60": "S-PCB-05",
    },
}

ROUTE_OPERATION_MATERIALS = {
    "RT-DJI-UAV-ASM-A01": {
        "100": {
            "M-BAT-01": 1,
            "M-MOT-02": 4,
            "M-PRP-03": 4,
            "M-FRM-04": 1,
            "S-PCB-05": 1,
            "S-OPT-06": 1,
        },
        "200": {
            "S-PCB-05": 1,
            "M-MOT-02": 4,
            "M-BAT-01": 1,
            "AUX-CAB-01": 1,
            "AUX-CAB-02": 4,
        },
        "300": {
            "M-BAT-01": 1,
            "S-PCB-05": 1,
        },
        "400": {
            "M-FRM-04": 1,
            "M-PRP-03": 4,
            "AUX-SCR-02": 8,
            "AUX-FIX-01": 1,
        },
        "500": {
            "S-OPT-06": 1,
        },
        "600": {
            "M-FRM-04": 1,
            "S-OPT-06": 1,
        },
        "800": {
            "AUX-CAB-04": 1,
        },
        "900": {
            "AUX-CAB-03": 1,
            "AUX-CON-01": 1,
            "AUX-CON-02": 1,
            "AUX-SIL-01": 1,
        },
    },
    "RT-DJI-PCB-A01": {
        "20": {
            "AUX-TAPE-01": 1,
        },
        "30": {
            "AUX-CLN-01": 1,
        },
        "40": {
            "AUX-CON-01": 1,
            "AUX-CON-02": 1,
            "AUX-CAB-01": 1,
            "AUX-CAB-02": 4,
        },
        "50": {
            "AUX-CLN-01": 1,
            "AUX-COT-02": 2,
        },
    },
}


def normalize(value):
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    return str(value).strip()


def load_rows(workbook_path: Path) -> dict[str, list[dict[str, str]]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        result: dict[str, list[dict[str, str]]] = {}
        for worksheet in workbook.worksheets:
            rows = list(worksheet.iter_rows(values_only=True))
            if not rows:
                result[worksheet.title] = []
                continue
            headers = [normalize(value) for value in rows[0]]
            parsed_rows: list[dict[str, str]] = []
            for raw_row in rows[1:]:
                row = {}
                has_value = False
                for index, header in enumerate(headers):
                    if not header:
                        continue
                    value = normalize(raw_row[index] if index < len(raw_row) else "")
                    if value:
                        has_value = True
                    row[header] = value
                if has_value:
                    parsed_rows.append(row)
            result[worksheet.title] = parsed_rows
        return result
    finally:
        workbook.close()


def feature_for_material(code: str, source: str, level: str) -> str:
    if level == "0":
        return "关键件"
    if source == "自制件":
        return "重要件"
    if any(token in code for token in ("BAT", "MOT", "OPT", "FRM")):
        return "重要件"
    return "一般件"


def serial_flag(code: str, level: str, source: str) -> str:
    if level == "0":
        return "是"
    if source == "自制件" and code.startswith("S-"):
        return "是"
    return "否"


def tooling_category(name: str, usage: str) -> str:
    if any(keyword in name for keyword in ("治具", "调试台", "调试器")):
        return "专用工装"
    if any(keyword in usage for keyword in ("调试", "校准", "检测")):
        return "专用工装"
    return "通用工具"


def parse_life_days(text: str) -> str:
    match = re.search(r"(\d+)\s*个月", text)
    if not match:
        return ""
    return str(int(match.group(1)) * 30)


def split_minutes(total_minutes: str) -> tuple[int, int]:
    total = int(float(total_minutes or "0"))
    return 0, total


def op_type(name: str) -> str:
    if any(keyword in name for keyword in ("检验", "测试", "校对")):
        return "检验"
    return "加工"


def route_content(op_row: dict[str, str]) -> str:
    parts = []
    involved = normalize(op_row.get("涉及物料"))
    if involved:
        parts.append(f"涉及物料：{involved}")
    control = normalize(op_row.get("关键控制点"))
    if control:
        parts.append(f"关键控制点：{control}")
    record = normalize(op_row.get("质量记录/防错")) or normalize(op_row.get("检测/记录要求"))
    if record:
        parts.append(f"记录要求：{record}")
    return "；".join(parts)


def root_material_code(bom_rows: list[dict[str, str]]) -> str:
    for row in bom_rows:
        if normalize(row.get("层级")) == "0":
            return normalize(row.get("子项物料编码"))
    raise ValueError("BOM 中未找到层级 0 的成品物料")


def build_admin_rows() -> list[dict[str, str]]:
    return [
        {"*父组织编码": "0", "行政组织类型": "公司", "*编码": ORG["admin_company"], "*名称": "智翔无人机科技有限公司", "简称": "智翔无人机"},
        {"*父组织编码": ORG["admin_company"], "行政组织类型": "部门", "*编码": ORG["admin_tech"], "*名称": "工艺技术部", "简称": "工艺技术"},
        {"*父组织编码": ORG["admin_company"], "行政组织类型": "部门", "*编码": ORG["admin_pmc"], "*名称": "生产计划部", "简称": "生产计划"},
        {"*父组织编码": ORG["admin_company"], "行政组织类型": "部门", "*编码": ORG["admin_qa"], "*名称": "质量管理部", "简称": "质量管理"},
        {"*父组织编码": ORG["admin_company"], "行政组织类型": "部门", "*编码": ORG["admin_wm"], "*名称": "仓储物流部", "简称": "仓储物流"},
        {"*父组织编码": ORG["admin_company"], "行政组织类型": "部门", "*编码": ORG["admin_em"], "*名称": "设备工装部", "简称": "设备工装"},
        {
            "*父组织编码": ORG["admin_company"],
            "行政组织类型": "工厂",
            "*编码": ORG["admin_plant"],
            "*名称": "大疆无人机",
            "简称": "大疆无人机",
            "备注": "负责无人机总装、PCB装联、调试测试与包装入库",
        },
    ]


def build_biz_rows() -> list[dict[str, str]]:
    rows = [
        {
            "*父组织编码": "0",
            "*密级": SECURITY,
            "*编码": ORG["biz_company"],
            "*名称": "智翔无人机科技有限公司",
            "简称": "智翔无人机",
            "工厂组织类型": "公司",
            "行政组织编码": ORG["admin_company"],
            "备注": "基于售前演示 BOM 与工艺路线生成的单工厂测试场景",
        },
        {
            "*父组织编码": ORG["biz_company"],
            "*密级": SECURITY,
            "*编码": ORG["biz_tech"],
            "*名称": "工艺技术部",
            "简称": "工艺技术",
            "工厂组织类型": "部门",
            "行政组织编码": ORG["admin_tech"],
        },
        {
            "*父组织编码": ORG["biz_company"],
            "*密级": SECURITY,
            "*编码": ORG["biz_pmc"],
            "*名称": "生产计划部",
            "简称": "生产计划",
            "工厂组织类型": "部门",
            "行政组织编码": ORG["admin_pmc"],
        },
        {
            "*父组织编码": ORG["biz_company"],
            "*密级": SECURITY,
            "*编码": ORG["biz_qa"],
            "*名称": "质量管理部",
            "简称": "质量管理",
            "工厂组织类型": "部门",
            "行政组织编码": ORG["admin_qa"],
        },
        {
            "*父组织编码": ORG["biz_company"],
            "*密级": SECURITY,
            "*编码": ORG["biz_wm"],
            "*名称": "仓储物流部",
            "简称": "仓储物流",
            "工厂组织类型": "部门",
            "行政组织编码": ORG["admin_wm"],
        },
        {
            "*父组织编码": ORG["biz_company"],
            "*密级": SECURITY,
            "*编码": ORG["biz_em"],
            "*名称": "设备工装部",
            "简称": "设备工装",
            "工厂组织类型": "部门",
            "行政组织编码": ORG["admin_em"],
        },
        {
            "*父组织编码": ORG["biz_company"],
            "*密级": SECURITY,
            "*编码": ORG["biz_plant"],
            "*名称": "大疆无人机",
            "简称": "大疆无人机",
            "工厂组织类型": "工厂",
            "行政组织编码": ORG["admin_plant"],
            "工厂类型": "装配专业",
            "备注": "单工厂无人机装配与电子装联演示场景",
        },
        {
            "*父组织编码": ORG["biz_plant"],
            "*密级": SECURITY,
            "*编码": ORG["biz_asm"],
            "*名称": "总装车间",
            "简称": "总装车间",
            "工厂组织类型": "车间",
            "行政组织编码": ORG["admin_plant"],
        },
        {
            "*父组织编码": ORG["biz_plant"],
            "*密级": SECURITY,
            "*编码": ORG["biz_smt"],
            "*名称": "SMT车间",
            "简称": "SMT车间",
            "工厂组织类型": "车间",
            "行政组织编码": ORG["admin_plant"],
        },
        {
            "*父组织编码": ORG["biz_plant"],
            "*密级": SECURITY,
            "*编码": ORG["biz_test"],
            "*名称": "调试测试车间",
            "简称": "调试测试",
            "工厂组织类型": "车间",
            "行政组织编码": ORG["admin_plant"],
        },
        {
            "*父组织编码": ORG["biz_plant"],
            "*密级": SECURITY,
            "*编码": ORG["biz_pack"],
            "*名称": "包装成品车间",
            "简称": "包装成品",
            "工厂组织类型": "车间",
            "行政组织编码": ORG["admin_plant"],
        },
    ]
    return rows


def build_user_rows() -> list[dict[str, str]]:
    rows = []
    for user_id, name, admin_code, biz_code, remark in USERS:
        rows.append(
            {
                "*编号": user_id,
                "名称": name,
                "用户安全等级": "一般",
                "性别": "男" if user_id in {"100", "103", "105", "106"} else "女",
                "行政组织编码": admin_code,
                "业务组织编码": biz_code,
                "备注": remark,
            }
        )
    return rows


def build_work_center_rows() -> list[dict[str, str]]:
    rows = []
    for name, wc in WORK_CENTERS.items():
        rows.append(
            {
                "*名称": name,
                "*类型": wc["type"],
                "*分类": wc["class"],
                "*编码": wc["code"],
                "*密级": SECURITY,
                "*工厂组织": wc["org"],
            }
        )
    return rows


def build_wc_user_rows() -> list[dict[str, str]]:
    return [{"*工作中心编码": code, "*用户": user_id} for code, user_id in WC_USER_LINKS.items()]


def build_warehouse_rows() -> list[dict[str, str]]:
    return [
        {
            "*工厂组织": ORG["biz_plant"],
            "*名称": "无人机原料库",
            "作业模式": "普通库房",
            "*业务类型": "ERP一级库",
            "*编码": "WH-DJI-RAW",
            "备注": "存放结构件、电子件和工装辅材",
            "*密级": SECURITY,
        },
        {
            "*工厂组织": ORG["biz_plant"],
            "*名称": "无人机线边库",
            "作业模式": "普通库房",
            "*业务类型": "车间二级库",
            "*编码": "WH-DJI-LINE",
            "备注": "为总装与SMT工位提供上线齐套和补料",
            "*密级": SECURITY,
        },
        {
            "*工厂组织": ORG["biz_plant"],
            "*名称": "无人机成品库",
            "作业模式": "普通库房",
            "*业务类型": "ERP二级库",
            "*编码": "WH-DJI-FG",
            "备注": "存放已检验放行的无人机成品",
            "*密级": SECURITY,
        },
    ]


def build_location_rows() -> list[dict[str, str]]:
    return [
        {
            "*工厂组织": ORG["biz_plant"],
            "*库房编码": "WH-DJI-RAW",
            "*名称": "结构与电子件区",
            "*编码": "LOC-DJI-RAW-01",
            "*密级": SECURITY,
            "备注": "存放机体、电机、PCB及光学模块",
        },
        {
            "*工厂组织": ORG["biz_plant"],
            "*库房编码": "WH-DJI-LINE",
            "*名称": "总装线边区",
            "*编码": "LOC-DJI-LINE-01",
            "*密级": SECURITY,
            "备注": "总装车间线边齐套与临时补料库位",
        },
        {
            "*工厂组织": ORG["biz_plant"],
            "*库房编码": "WH-DJI-FG",
            "*名称": "成品待发区",
            "*编码": "LOC-DJI-FG-01",
            "*密级": SECURITY,
            "备注": "包装完成后的无人机成品暂存区",
        },
    ]


def build_material_rows(
    bom_rows: list[dict[str, str]],
    aux_rows: list[dict[str, str]],
) -> tuple[list[dict[str, str]], dict[str, dict[str, str]], dict[str, str]]:
    rows: list[dict[str, str]] = []
    material_by_code: dict[str, dict[str, str]] = {}
    name_to_code: dict[str, str] = {}

    for bom_row in bom_rows:
        code = normalize(bom_row["子项物料编码"])
        name = normalize(bom_row["子项物料名称"])
        make_type = "自制件" if normalize(bom_row["来源"]) == "自制" else "外购件"
        unit = normalize(bom_row["单位"]) or "个"
        remark = normalize(bom_row.get("备注"))
        if unit != "个":
            remark = f"{remark}；原始单位：{unit}" if remark else f"原始单位：{unit}"
        row = {
            "*物料分类": "物料",
            "*名称": name,
            "物料类别": "零部件",
            "图号": code,
            "*制造类型": make_type,
            "计量单位": "个",
            "特性分类": feature_for_material(code, make_type, normalize(bom_row["层级"])),
            "启用批次标记": "是",
            "启用序列号标记": serial_flag(code, normalize(bom_row["层级"]), make_type),
            "物料阶段": "量产",
            "发布版本时间": RELEASE_TIME,
            "发布人": "100",
            "*版本号": VERSION,
            "*编码": code,
            "*密级": SECURITY,
            "备注": remark,
        }
        rows.append(row)
        material_by_code[code] = row
        name_to_code[name] = code

    for aux_row in aux_rows:
        code = normalize(aux_row["物料编码"])
        name = normalize(aux_row["物料名称"])
        remark_parts = [
            f"适用工序：{normalize(aux_row.get('适用工序'))}" if normalize(aux_row.get("适用工序")) else "",
            f"有效期：{normalize(aux_row.get('有效期'))}" if normalize(aux_row.get("有效期")) else "",
            normalize(aux_row.get("备注")),
        ]
        row = {
            "*物料分类": "物料",
            "*名称": name,
            "物料类别": "辅助材料",
            "图号": code,
            "规格": normalize(aux_row.get("规格 / 型号")),
            "*制造类型": "外购件",
            "计量单位": "个",
            "特性分类": "一般件",
            "启用批次标记": "是",
            "启用序列号标记": "否",
            "物料阶段": "量产",
            "发布版本时间": RELEASE_TIME,
            "发布人": "100",
            "*版本号": VERSION,
            "*编码": code,
            "*密级": SECURITY,
            "备注": "；".join(part for part in remark_parts if part),
        }
        rows.append(row)
        material_by_code[code] = row
        name_to_code[name] = code

    name_to_code["PCB"] = "S-PCB-05"
    name_to_code["主板"] = "S-PCB-05"
    name_to_code["光模块"] = "S-OPT-06"
    name_to_code["光学模块"] = "S-OPT-06"
    name_to_code["整机"] = "DJI-UAV-01"
    name_to_code["全部物料"] = "DJI-UAV-01"
    return rows, material_by_code, name_to_code


def build_tooling_rows(tool_rows: list[dict[str, str]]) -> list[dict[str, str]]:
    rows = []
    for tool_row in tool_rows:
        code = normalize(tool_row["工具编码"])
        name = normalize(tool_row["工具名称"])
        usage = normalize(tool_row.get("用途"))
        precision = normalize(tool_row.get("精度要求"))
        cycle = normalize(tool_row.get("定检周期"))
        last_check = normalize(tool_row.get("最近检定日期（示例）"))
        manage = normalize(tool_row.get("管理要求"))
        remark_parts = []
        if usage:
            remark_parts.append(f"用途：{usage}")
        if precision:
            remark_parts.append(f"精度要求：{precision}")
        if cycle:
            remark_parts.append(f"定检周期：{cycle}")
        if last_check and last_check not in {"-", "/"}:
            remark_parts.append(f"最近检定日期：{last_check}")
        if manage:
            remark_parts.append(f"管理要求：{manage}")
        rows.append(
            {
                "*物料分类": "工装工具",
                "*名称": name,
                "物料类别": "辅助材料",
                "图号": code,
                "型号": normalize(tool_row.get("规格/型号")),
                "规格": usage,
                "*制造类型": "外购件",
                "计量单位": "个",
                "特性分类": "重要件" if precision and precision != "/" else "一般件",
                "启用批次标记": "否",
                "启用序列号标记": "否",
                "工装类别": tooling_category(name, usage),
                "理论寿命(天)": parse_life_days(cycle),
                "一次性工装标记": "否",
                "单件工装标记": "否",
                "物料阶段": "量产",
                "发布版本时间": RELEASE_TIME,
                "发布人": "100",
                "*版本号": VERSION,
                "*编码": code,
                "备注": "；".join(remark_parts),
                "*密级": SECURITY,
            }
        )
    return rows


def build_mbom_rows(
    bom_rows: list[dict[str, str]],
    material_by_code: dict[str, dict[str, str]],
    root_code: str,
) -> tuple[list[dict[str, str]], list[dict[str, str]], str]:
    mbom_code = "MBOM-DJI-UAV-A01"
    mbom_rows = [
        {
            "*物料版本号": VERSION,
            "*物料编码": root_code,
            "*版本号": VERSION,
            "*编码": mbom_code,
            "*密级": SECURITY,
            "名称": "大疆无人机MBOM",
            "发布人": "100",
        }
    ]
    node_rows = []
    for index, bom_row in enumerate(bom_rows, start=1):
        code = normalize(bom_row["子项物料编码"])
        material_row = material_by_code[code]
        node = {
            "*MBOM版本号": VERSION,
            "*MBOM编码": mbom_code,
            "*物料编码": code,
            "物料名称": normalize(bom_row["子项物料名称"]),
            "物料图号": code,
            "*物料版本": VERSION,
            "*物料类别": normalize(material_row["物料类别"]),
            "制造类型": normalize(material_row["*制造类型"]),
            "数量": float(bom_row["数量"]) if "." in normalize(bom_row["数量"]) else int(float(bom_row["数量"])),
            "计量单位": "个",
            "*层级": int(float(normalize(bom_row["层级"]))),
            "*序号": index * 10,
            "物料阶段": "量产",
        }
        parent_code = normalize(bom_row.get("父项物料编码"))
        if parent_code and parent_code != "-":
            node["父物料编码"] = parent_code
            node["父物料版本"] = VERSION
        node_rows.append(node)
    return mbom_rows, node_rows, mbom_code


def build_process_lib_rows(route_sheet_name: str, route_rows: list[dict[str, str]]) -> list[dict[str, str]]:
    route_config = ROUTES[route_sheet_name]
    result = []
    for op_row in route_rows:
        prep, run = split_minutes(normalize(op_row["工时（分钟）"]))
        wc = WORK_CENTERS[normalize(op_row["工作中心"])]
        result.append(
            {
                "序专业类型": route_config["op_spec"],
                "*名称": normalize(op_row["工序名称"]),
                "*工序类型": op_type(normalize(op_row["工序名称"])),
                "*工作中心编码": wc["code"],
                "*定额准备时间": prep,
                "*定额加工时间": run,
                "执行标记": "是",
                "*时间单位": "分钟",
                "产出比": 1,
                "工序内容": route_content(op_row),
                "*密级": SECURITY,
            }
        )
    return result


def build_route_rows(
    route_sheet_name: str,
    route_source_rows: list[dict[str, str]],
    material_by_code: dict[str, dict[str, str]],
) -> tuple[list[dict[str, str]], list[dict[str, str]], list[dict[str, str]], list[dict[str, str]], list[dict[str, str]]]:
    route_config = ROUTES[route_sheet_name]
    route_code = route_config["code"]
    route_rows = [
        {
            "*名称": route_config["name"],
            "*工艺类型": "正式工艺",
            "工艺专业": route_config["route_spec"],
            "物料版本号": VERSION,
            "物料编码": route_config["material_code"],
            "*版本号": VERSION,
            "*编码": route_code,
            "*密级": SECURITY,
            "*工厂组织": route_config["org"],
            "备注": f"由附件《{route_sheet_name}》自动转换生成",
            "发布人": "100",
        }
    ]
    op_rows = []
    seq_rows = []
    op_material_rows = []
    step_rows = []
    previous_no = ""

    for op_row in route_source_rows:
        op_no = normalize(op_row["工序序号"])
        op_name_text = normalize(op_row["工序名称"])
        prep, run = split_minutes(normalize(op_row["工时（分钟）"]))
        wc_name = normalize(op_row["工作中心"])
        wc = WORK_CENTERS[wc_name]
        output_material_code = ROUTE_OUTPUTS.get(route_code, {}).get(op_no, route_config["material_code"])
        record_text = normalize(op_row.get("质量记录/防错")) or normalize(op_row.get("检测/记录要求"))

        route_op = {
            "*工序号": op_no,
            "*工序类型": op_type(op_name_text),
            "工序内容": route_content(op_row),
            "*工作中心编码": wc["code"],
            "*定额辅助工时": prep,
            "*定额加工时间": run,
            "*时间单位": "分钟",
            "执行标记": "是",
            "产出比": 1,
            "*工艺路线版本号": VERSION,
            "*工艺路线编码": route_code,
            "*工序名称": op_name_text,
            "产出物料版本号": VERSION,
            "产出物料编码": output_material_code,
            "工序专业类型": route_config["op_spec"],
        }
        if previous_no:
            route_op["前置工序"] = previous_no
        op_rows.append(route_op)

        if previous_no:
            seq_rows.append(
                {
                    "*接续关系": "ES",
                    "*工序号": op_no,
                    "*上道工序号": previous_no,
                    "*工艺路线版本号": VERSION,
                    "*工艺路线编码": route_code,
                }
            )

        material_map = ROUTE_OPERATION_MATERIALS.get(route_code, {}).get(op_no, {})
        for material_code, qty in material_map.items():
            if material_code not in material_by_code:
                continue
            op_material_rows.append(
                {
                    "*工艺路线版本号": VERSION,
                    "*工艺路线编码": route_code,
                    "*物料版本号": VERSION,
                    "*物料编码": material_code,
                    "*工序号": op_no,
                    "数量": qty,
                }
            )

        key_control = normalize(op_row.get("关键控制点"))
        if key_control:
            step_rows.append(
                {
                    "*工艺路线版本号": VERSION,
                    "*工艺路线编码": route_code,
                    "*工序号": op_no,
                    "*工步序号": 10,
                    "*工步名称": "关键控制",
                    "工步内容": key_control,
                }
            )
        if record_text:
            step_rows.append(
                {
                    "*工艺路线版本号": VERSION,
                    "*工艺路线编码": route_code,
                    "*工序号": op_no,
                    "*工步序号": 20,
                    "*工步名称": "质量记录",
                    "工步内容": record_text,
                }
            )

        previous_no = op_no

    return route_rows, op_rows, seq_rows, op_material_rows, step_rows


def build_orders(
    material_by_code: dict[str, dict[str, str]],
    route_op_rows: list[dict[str, str]],
    route_material_rows: list[dict[str, str]],
    mbom_code: str,
) -> tuple[list[dict[str, str]], list[dict[str, str]]]:
    orders = [
        {
            "订单类型": "标准",
            "*物料版本号": VERSION,
            "*物料编码": "S-PCB-05",
            "BOM版本号": "",
            "BOM编码": "",
            "制造型号": "S-PCB-05",
            "工艺路线版本号": VERSION,
            "工艺路线编码": "RT-DJI-PCB-A01",
            "*计量单位": "个",
            "*计划数量": 6,
            "计划产出数量": 6,
            "合格数量": 0,
            "报废数量": 0,
            "已释放数量": 0,
            "*计划开始时间": "2026-05-18 08:00:00",
            "*计划结束时间": "2026-05-18 14:00:00",
            "计划类型": "零部件加工计划",
            "排产状态": "无",
            "业务状态": "初始",
            "优先级": 5,
            "计划员": "101",
            "*控制状态": "正常",
            "释放状态": "未释放",
            "*所属组织": ORG["biz_smt"],
            "*编码": "MO-DJI-UAV-001",
            "备注": "主控印制电路板演示生产订单",
            "*密级": SECURITY,
            "集成系统": INTEGRATION_SYSTEM,
            "集成数据主键": "MO-DJI-UAV-001",
            "*集成创建时间": "2026-05-17 16:00:00",
        },
        {
            "订单类型": "标准",
            "*物料版本号": VERSION,
            "*物料编码": "DJI-UAV-01",
            "BOM版本号": VERSION,
            "BOM编码": mbom_code,
            "制造型号": "DJI-UAV-01",
            "工艺路线版本号": VERSION,
            "工艺路线编码": "RT-DJI-UAV-ASM-A01",
            "*计量单位": "个",
            "*计划数量": 3,
            "计划产出数量": 3,
            "合格数量": 0,
            "报废数量": 0,
            "已释放数量": 0,
            "*计划开始时间": "2026-05-19 08:00:00",
            "*计划结束时间": "2026-05-19 18:00:00",
            "计划类型": "零部件加工计划",
            "排产状态": "无",
            "业务状态": "初始",
            "优先级": 4,
            "计划员": "101",
            "*控制状态": "正常",
            "释放状态": "未释放",
            "*所属组织": ORG["biz_asm"],
            "*编码": "MO-DJI-UAV-002",
            "备注": "大疆无人机整机演示生产订单",
            "*密级": SECURITY,
            "集成系统": INTEGRATION_SYSTEM,
            "集成数据主键": "MO-DJI-UAV-002",
            "*集成创建时间": "2026-05-18 16:00:00",
        },
    ]

    order_qty = {row["*编码"]: int(row["*计划数量"]) for row in orders}
    route_to_order = {
        "RT-DJI-PCB-A01": "MO-DJI-UAV-001",
        "RT-DJI-UAV-ASM-A01": "MO-DJI-UAV-002",
    }
    op_name_map = {
        (normalize(row["*工艺路线编码"]), normalize(row["*工序号"])): normalize(row["*工序名称"])
        for row in route_op_rows
    }
    pick_rows = []
    for row in route_material_rows:
        route_code = normalize(row["*工艺路线编码"])
        order_code = route_to_order.get(route_code)
        if not order_code:
            continue
        plan_qty = order_qty[order_code]
        unit_qty = float(row["数量"])
        demand_qty = int(unit_qty * plan_qty) if unit_qty.is_integer() else round(unit_qty * plan_qty, 3)
        op_no = normalize(row["*工序号"])
        pick_rows.append(
            {
                "*生产订单编码": order_code,
                "*物料版本号": VERSION,
                "*物料编码": normalize(row["*物料编码"]),
                "工序名称": op_name_map.get((route_code, op_no), ""),
                "*需求数量": demand_qty,
                "*子件比例": int(unit_qty) if unit_qty.is_integer() else unit_qty,
                "替换件物料版本号": "",
                "替换件物料编码": "",
                "是否必须装入": "是",
                "工序编码": op_no,
            }
        )
    return orders, pick_rows


def build_seed(source_rows: dict[str, list[dict[str, str]]]) -> dict:
    bom_rows = source_rows["大疆无人机"]
    uav_route_rows = source_rows["无人机总装工艺路线"]
    pcb_route_rows = source_rows["电控印制板工艺路线"]
    aux_rows = source_rows["辅料信息表"]
    tool_rows = source_rows["工具工装台帐表"]

    root_code = root_material_code(bom_rows)
    material_rows, material_by_code, _name_to_code = build_material_rows(bom_rows, aux_rows)
    tooling_rows = build_tooling_rows(tool_rows)
    mbom_rows, mbom_node_rows, mbom_code = build_mbom_rows(bom_rows, material_by_code, root_code)
    process_lib_rows = build_process_lib_rows("无人机总装工艺路线", uav_route_rows)
    process_lib_rows.extend(build_process_lib_rows("电控印制板工艺路线", pcb_route_rows))

    route_rows = []
    route_op_rows = []
    route_seq_rows = []
    route_material_rows = []
    route_step_rows = []

    for sheet_name, rows in (
        ("无人机总装工艺路线", uav_route_rows),
        ("电控印制板工艺路线", pcb_route_rows),
    ):
        built = build_route_rows(sheet_name, rows, material_by_code)
        route_rows.extend(built[0])
        route_op_rows.extend(built[1])
        route_seq_rows.extend(built[2])
        route_material_rows.extend(built[3])
        route_step_rows.extend(built[4])

    order_rows, pick_rows = build_orders(material_by_code, route_op_rows, route_material_rows, mbom_code)

    return {
        "metadata": {
            "name": "大疆无人机售前演示MOM种子",
            "industry": "无人机制造",
            "product_family": "多旋翼无人机",
            "product_model": "DJI-UAV-01 大疆无人机",
            "description": "根据售前提供的 BOM、总装工艺路线、PCB工艺路线、辅料与工具台账整理出的单工厂测试导入数据。",
            "default_version": VERSION,
            "default_security": SECURITY,
            "volume_profile": "演示版",
            "namespace": "DJIUAV01",
        },
        "external_references": ["0"],
        "workbooks": {
            "系统配置_模板.xlsx": {
                "行政组织": build_admin_rows(),
                "业务组织": build_biz_rows(),
                "用户": build_user_rows(),
            },
            "工厂资源_模板.xlsx": {
                "供应商": [],
                "设备": [],
                "设备与用户的关系实体类": [],
                "工装工具": tooling_rows,
                "工装检定策略关系": [],
                "工装保养策略关系": [],
                "工作中心": build_work_center_rows(),
                "工作中心与用户关系": build_wc_user_rows(),
                "工作中心与供应商的关系": [],
                "工作中心与设备的关系": [],
                "库房": build_warehouse_rows(),
                "库位": build_location_rows(),
                "工序库": process_lib_rows,
            },
            "产品与工艺_模板.xlsx": {
                "物料": material_rows,
                "MBOM": mbom_rows,
                "MBOM节点": mbom_node_rows,
                "工艺路线": route_rows,
                "工艺路线工序": route_op_rows,
                "工艺路线工序序列": route_seq_rows,
                "工艺路线工序物料": route_material_rows,
                "工艺路线工步": route_step_rows,
            },
            "生产订单_模板.xlsx": {
                "生产订单": order_rows,
                "备料清单": pick_rows,
            },
        },
    }


def resolve_skill_dir(explicit_path: Path | None) -> Path:
    candidates = []
    if explicit_path is not None:
        candidates.append(explicit_path)
    codex_home = os.environ.get("CODEX_HOME")
    if codex_home:
        candidates.append(Path(codex_home) / "skills" / "mom-business-data-generator")
    candidates.append(Path.home() / ".codex" / "skills" / "mom-business-data-generator")
    for candidate in candidates:
        if candidate and (candidate / "scripts" / "generate_seed_workbooks.py").exists():
            return candidate
    raise FileNotFoundError("未找到 mom-business-data-generator skill 目录，请使用 --skill-dir 指定")


def load_skill_generator(skill_dir: Path):
    script_path = skill_dir / "scripts" / "generate_seed_workbooks.py"
    spec = importlib.util.spec_from_file_location("mom_seed_generator", script_path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"无法加载 skill 脚本: {script_path}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def write_package(seed: dict, skill_dir: Path, package_dir: Path, package_zip: Path):
    generator = load_skill_generator(skill_dir)
    template_dir = skill_dir / "templates"
    headers = generator.load_template_headers(template_dir)
    validations = generator.load_template_list_validations(template_dir)
    structure_errors, warnings = generator.validate_structure(seed, headers)
    errors = (
        structure_errors
        + generator.validate_list_values(seed, validations)
        + generator.validate_unique(seed)
        + generator.validate_references(seed)
        + generator.validate_storage_org_rules(seed)
        + generator.validate_route_refs(seed)
        + generator.validate_order_refs(seed)
    )
    for warning in warnings:
        print(f"[WARN] {warning}")
    if errors:
        print("[ERROR] 生成前校验失败：", file=sys.stderr)
        for error in errors:
            print(f"  - {error}", file=sys.stderr)
        raise SystemExit(1)

    summary = generator.build_summary(seed)
    package_dir.mkdir(parents=True, exist_ok=True)
    for workbook_name in generator.TEMPLATE_FILES:
        generator.fill_workbook(
            template_dir / workbook_name,
            package_dir / workbook_name,
            seed.get("workbooks", {}).get(workbook_name, {}),
        )
        print(f"[OK] 已生成 {package_dir / workbook_name}")
    summary_path = package_dir / "种子概览.json"
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"[OK] 已生成 {summary_path}")

    package_zip.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(package_zip, "w", zipfile.ZIP_DEFLATED) as archive:
        for file_path in sorted(package_dir.iterdir()):
            if file_path.is_file():
                archive.write(file_path, arcname=file_path.name)
    print(f"[OK] 已生成 {package_zip}")


def write_route_only_workbook(seed: dict, skill_dir: Path, output_path: Path):
    generator = load_skill_generator(skill_dir)
    template_path = skill_dir / "templates" / "产品与工艺_模板.xlsx"
    route_seed = {
        "工艺路线": seed.get("workbooks", {}).get("产品与工艺_模板.xlsx", {}).get("工艺路线", []),
        "工艺路线工序": seed.get("workbooks", {}).get("产品与工艺_模板.xlsx", {}).get("工艺路线工序", []),
        "工艺路线工序序列": seed.get("workbooks", {}).get("产品与工艺_模板.xlsx", {}).get("工艺路线工序序列", []),
        "工艺路线工序物料": seed.get("workbooks", {}).get("产品与工艺_模板.xlsx", {}).get("工艺路线工序物料", []),
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    generator.fill_workbook(template_path, output_path, route_seed)

    workbook = load_workbook(output_path)
    for sheet_name in list(workbook.sheetnames):
        if sheet_name not in ROUTE_ONLY_SHEETS:
            del workbook[sheet_name]
    workbook.save(output_path)
    print(f"[OK] 已生成工艺路线单独导入: {output_path}")


def parse_args():
    parser = argparse.ArgumentParser(description="根据附件 Excel 生成 DJI 无人机 MOM 导入种子与导入包")
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT, help="售前 Excel 输入文件")
    parser.add_argument("--seed-output", type=Path, default=DEFAULT_SEED, help="生成的 seed JSON 路径")
    parser.add_argument("--package-dir", type=Path, default=DEFAULT_PACKAGE_DIR, help="导入 Excel 包输出目录")
    parser.add_argument("--package-zip", type=Path, default=DEFAULT_PACKAGE_ZIP, help="导入包 zip 输出路径")
    parser.add_argument(
        "--route-only-workbook",
        type=Path,
        default=DEFAULT_ROUTE_ONLY_WORKBOOK,
        help="单独输出的工艺路线导入工作簿路径",
    )
    parser.add_argument("--skill-dir", type=Path, default=None, help="mom-business-data-generator skill 目录")
    parser.add_argument("--seed-only", action="store_true", help="仅生成 seed JSON，不生成导入 Excel 包")
    return parser.parse_args()


def main():
    args = parse_args()
    input_path = args.input.resolve()
    if not input_path.exists():
        print(f"输入文件不存在: {input_path}", file=sys.stderr)
        return 1

    source_rows = load_rows(input_path)
    seed = build_seed(source_rows)

    args.seed_output.parent.mkdir(parents=True, exist_ok=True)
    args.seed_output.write_text(json.dumps(seed, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"[OK] 已生成种子文件: {args.seed_output.resolve()}")

    if args.seed_only:
        return 0

    skill_dir = resolve_skill_dir(args.skill_dir)
    write_package(seed, skill_dir, args.package_dir.resolve(), args.package_zip.resolve())
    write_route_only_workbook(seed, skill_dir, args.route_only_workbook.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
