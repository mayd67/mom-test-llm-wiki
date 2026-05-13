from __future__ import annotations

import argparse
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass


PLAN_CODES = {
    "RtDJ001": ("QS-DJ30S01-RT001", "无人机总装过程检方案"),
    "RtDJ002": ("QS-DJ30S01-RT002", "电控印制板过程检方案"),
}


EXPLICIT_ITEMS: dict[tuple[str, str], list[dict[str, str]]] = {
    ("RtDJ001", "100"): [
        {"name": "物料齐套状态", "tool": "", "feature": "定性", "std": "齐套", "upper": "", "lower": "", "method": "目视+清单核对", "note": "确认来料、版本、数量和批次齐套。"},
        {"name": "防静电测试阻值", "tool_name": "手腕带测试仪", "feature": "定量", "std": "5MΩ", "upper": "10", "lower": "0.75", "method": "仪表检测", "note": "上岗前确认ESD状态满足要求。"},
    ],
    ("RtDJ001", "200"): [
        {"name": "线束接插正确性", "tool": "", "feature": "定性", "std": "合格", "upper": "", "lower": "", "method": "目视", "note": "核对线束型号、接口方向和锁止状态。"},
        {"name": "关键锁付扭矩", "tool_name": "数显扭力螺丝刀", "feature": "定量", "std": "0.45N·m", "upper": "0.50", "lower": "0.40", "method": "力矩检测", "note": "适用于PCB、电装小螺钉紧固。"},
    ],
    ("RtDJ001", "300"): [
        {"name": "上电输入电压", "tool_name": "数字万用表", "feature": "定量", "std": "16V", "upper": "16.8", "lower": "15.2", "method": "仪表检测", "note": "首上电前确认供电电压范围。"},
        {"name": "极性连接正确性", "tool_name": "数字万用表", "feature": "定性", "std": "合格", "upper": "", "lower": "", "method": "通电确认", "note": "确认正负极连接无反接。"},
    ],
    ("RtDJ001", "400"): [
        {"name": "电机安装扭矩", "tool_name": "扭矩扳手", "feature": "定量", "std": "1.20N·m", "upper": "1.30", "lower": "1.10", "method": "力矩检测", "note": "适用于电机与机体连接螺钉。"},
        {"name": "螺旋桨动平衡", "tool_name": "螺旋桨动平衡仪", "feature": "定量", "std": "0.10g·mm", "upper": "0.15", "lower": "0", "method": "平衡检测", "note": "确保螺旋桨总成动平衡在控制范围。"},
    ],
    ("RtDJ001", "500"): [
        {"name": "零位偏差", "tool_name": "电位器调试工装", "feature": "定量", "std": "0%", "upper": "0.5", "lower": "-0.5", "method": "工装校准", "note": "电位器零位校准偏差。"},
        {"name": "曲线连续性", "tool_name": "飞控调试器", "feature": "定性", "std": "合格", "upper": "", "lower": "", "method": "软件判定", "note": "曲线输出连续、无跳变。"},
    ],
    ("RtDJ001", "600"): [
        {"name": "平台水平角", "tool_name": "云台稳定调试台", "feature": "定量", "std": "0°", "upper": "0.10", "lower": "-0.10", "method": "姿态检测", "note": "静态水平角偏差。"},
        {"name": "姿态自稳响应", "tool_name": "云台稳定调试台", "feature": "定性", "std": "合格", "upper": "", "lower": "", "method": "动作确认", "note": "IMU/云台自稳响应及时，无抖动。"},
    ],
    ("RtDJ001", "700"): [
        {"name": "基础功能自检", "tool_name": "飞控调试器", "feature": "定性", "std": "通过", "upper": "", "lower": "", "method": "软件检测", "note": "电源、飞控、姿态、通信基础功能验证。"},
        {"name": "空载工作电流", "tool_name": "电流钳/功率计", "feature": "定量", "std": "1.8A", "upper": "2.2", "lower": "1.4", "method": "电流检测", "note": "整机初次加电空载电流值。"},
    ],
    ("RtDJ001", "800"): [
        {"name": "线性响应偏差", "tool_name": "飞控调试器", "feature": "定量", "std": "0%", "upper": "1.0", "lower": "-1.0", "method": "软件校准", "note": "控制输入与反馈输出线性度。"},
        {"name": "校准结果", "tool_name": "飞控调试器", "feature": "定性", "std": "合格", "upper": "", "lower": "", "method": "软件判定", "note": "传感器线性校准通过。"},
    ],
    ("RtDJ001", "900"): [
        {"name": "连接器闭锁状态", "tool": "", "feature": "定性", "std": "合格", "upper": "", "lower": "", "method": "目视", "note": "确认连接器锁止、排线压接和布线闭环。"},
        {"name": "三防胶覆盖完整性", "tool": "", "feature": "定性", "std": "合格", "upper": "", "lower": "", "method": "目视", "note": "关键区域三防胶喷涂完整、无漏喷。"},
    ],
    ("RtDJ001", "1000"): [
        {"name": "稳态工作电流", "tool_name": "电流钳/功率计", "feature": "定量", "std": "3.2A", "upper": "3.8", "lower": "2.6", "method": "电流检测", "note": "整机稳定运行时的工作电流。"},
        {"name": "功率曲线稳定性", "tool_name": "电流钳/功率计", "feature": "定性", "std": "合格", "upper": "", "lower": "", "method": "曲线判定", "note": "功率曲线平稳，无突变。"},
    ],
    ("RtDJ001", "1100"): [
        {"name": "整机外观", "tool_name": "放大镜/显微镜", "feature": "定性", "std": "合格", "upper": "", "lower": "", "method": "目视", "note": "外观无磕碰、划伤、污渍和明显装配缺陷。"},
        {"name": "标签信息正确性", "tool": "", "feature": "定性", "std": "合格", "upper": "", "lower": "", "method": "目视+记录复核", "note": "标签内容、位置和序列号一致。"},
    ],
    ("RtDJ001", "1200"): [
        {"name": "包装BOM符合性", "tool": "", "feature": "定性", "std": "合格", "upper": "", "lower": "", "method": "清单核对", "note": "整机与附件按项目包装清单齐套。"},
        {"name": "包装密封状态", "tool": "", "feature": "定性", "std": "合格", "upper": "", "lower": "", "method": "目视", "note": "包装封装完整，防护材料到位。"},
    ],
    ("RtDJ002", "10"): [
        {"name": "料站核对一致性", "tool": "", "feature": "定性", "std": "合格", "upper": "", "lower": "", "method": "清单核对", "note": "贴片料站与工单程序一致。"},
        {"name": "MSL烘烤记录完整性", "tool": "", "feature": "定性", "std": "合格", "upper": "", "lower": "", "method": "记录复核", "note": "湿敏器件烘烤记录齐全、状态有效。"},
    ],
    ("RtDJ002", "20"): [
        {"name": "炉温峰值", "tool_name": "回流焊炉", "feature": "定量", "std": "245℃", "upper": "250", "lower": "240", "method": "曲线检测", "note": "回流焊峰值温度控制。"},
        {"name": "贴片偏移量", "tool_name": "AOI光学检测仪", "feature": "定量", "std": "0mm", "upper": "0.05", "lower": "-0.05", "method": "AOI检测", "note": "关键元件贴装偏移量。"},
    ],
    ("RtDJ002", "30"): [
        {"name": "表面洁净度", "tool_name": "放大镜/显微镜", "feature": "定性", "std": "合格", "upper": "", "lower": "", "method": "目视", "note": "清洗后PCB表面洁净，无明显残留。"},
        {"name": "助焊剂残留状态", "tool_name": "放大镜/显微镜", "feature": "定性", "std": "合格", "upper": "", "lower": "", "method": "目视", "note": "无肉眼可见助焊剂污染。"},
    ],
    ("RtDJ002", "40"): [
        {"name": "烙铁头温度", "tool_name": "防静电电烙铁", "feature": "定量", "std": "350℃", "upper": "365", "lower": "335", "method": "仪表检测", "note": "手工焊接温度点检。"},
        {"name": "焊点外观", "tool_name": "放大镜/显微镜", "feature": "定性", "std": "合格", "upper": "", "lower": "", "method": "目视", "note": "焊点饱满，无连焊、虚焊、冷焊。"},
    ],
    ("RtDJ002", "50"): [
        {"name": "局部残留清洁度", "tool_name": "放大镜/显微镜", "feature": "定性", "std": "合格", "upper": "", "lower": "", "method": "目视", "note": "局部焊接区域无残胶、无残渣。"},
        {"name": "接插件清洁状态", "tool_name": "放大镜/显微镜", "feature": "定性", "std": "合格", "upper": "", "lower": "", "method": "目视", "note": "连接器周边清洁，无异物污染。"},
    ],
    ("RtDJ002", "60"): [
        {"name": "ICT通断测试", "tool_name": "ICT在线测试仪", "feature": "定性", "std": "通过", "upper": "", "lower": "", "method": "在线检测", "note": "PCB关键网络、元件值检测通过。"},
        {"name": "外观检验", "tool_name": "AOI光学检测仪", "feature": "定性", "std": "合格", "upper": "", "lower": "", "method": "AOI+目视", "note": "外观、器件极性、焊点质量符合要求。"},
    ],
}


def normalize(value) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    return str(value).strip()


def table_rows(ws) -> list[list[str]]:
    rows: list[list[str]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        values = [normalize(cell) for cell in row]
        if any(values):
            rows.append(values)
    return rows


def load_route_data(route_workbook: Path) -> tuple[list[dict], dict[str, tuple[str, str]], dict[str, str]]:
    wb = load_workbook(route_workbook, data_only=True)
    try:
        material_rows = table_rows(wb.worksheets[1])
        route_rows = table_rows(wb.worksheets[5])
        op_rows = table_rows(wb.worksheets[6])

        material_name = {row[24]: row[1] for row in material_rows if len(row) > 24 and row[24]}
        routes = []
        op_map: dict[str, tuple[str, str]] = {}
        for row in route_rows:
            route_code = row[8]
            routes.append(
                {
                    "route_name": row[0],
                    "route_type": row[1],
                    "material_code": row[4],
                    "material_name": material_name.get(row[4], ""),
                    "version": row[7],
                    "route_code": route_code,
                    "factory_org": row[11],
                }
            )
        for row in op_rows:
            route_code = row[13]
            op_map[f"{route_code}:{row[0]}"] = (row[14], row[2])
        return routes, op_map, material_name
    finally:
        wb.close()


def load_tool_map(resource_workbook: Path) -> dict[str, str]:
    wb = load_workbook(resource_workbook, data_only=True)
    try:
        tool_rows = table_rows(wb.worksheets[4])
        tool_map: dict[str, str] = {}
        for row in tool_rows:
            name = row[1]
            version = row[23]
            code = row[24]
            if name and code:
                tool_map[name] = f"{name} | {version or 'A.01'} | {code}"
        return tool_map
    finally:
        wb.close()


def fallback_items(route_code: str, op_no: str, op_name: str, op_content: str) -> list[dict[str, str]]:
    return [
        {
            "name": f"{op_name}结果确认",
            "tool": "",
            "feature": "定性",
            "std": "合格",
            "upper": "",
            "lower": "",
            "method": "目视+记录复核",
            "note": op_content or f"{route_code}-{op_no} 默认检查项",
        }
    ]


def build_workbook(route_workbook: Path, resource_workbook: Path, output_path: Path) -> Path:
    routes, op_map, material_name = load_route_data(route_workbook)
    tool_map = load_tool_map(resource_workbook)

    wb = Workbook()
    ws_plan = wb.active
    ws_plan.title = "质量方案"
    ws_plan.append(["方案编码", "方案名称", "方案类型", "工艺路线编码", "工艺路线名称", "物料编码", "物料名称", "状态", "备注"])

    ws_items = wb.create_sheet("检查项配置")
    ws_items.append(
        [
            "方案编码",
            "工艺路线编码",
            "工艺路线名称",
            "工序号",
            "工序名称",
            "序号",
            "编码",
            "名称",
            "工装",
            "检验特征",
            "标准值",
            "上限值",
            "下限值",
            "检验方法",
            "说明",
        ]
    )

    ws_readme = wb.create_sheet("说明")
    ws_readme.append(["说明"])
    ws_readme.append(["1. 本文件根据最新工艺路线结果自动生成，字段结构对齐质量方案管理-检查项配置页面。"])
    ws_readme.append(["2. 方案类型统一为“过程检验”，检查项按工艺路线-工序粒度展开。"])
    ws_readme.append(["3. 工装列优先引用现有工装工具编码，格式为“名称 | 版本 | 编码”。"])

    for route in routes:
        route_code = route["route_code"]
        plan_code, plan_name = PLAN_CODES.get(route_code, (f"QS-{route_code}", f"{route['route_name']}过程检方案"))
        ws_plan.append(
            [
                plan_code,
                plan_name,
                "过程检验",
                route_code,
                route["route_name"],
                route["material_code"],
                route["material_name"],
                "启用",
                f"根据工艺路线 {route['route_name']} 自动生成",
            ]
        )

        route_items = sorted(
            [(key, value) for key, value in EXPLICIT_ITEMS.items() if key[0] == route_code],
            key=lambda item: int(item[0][1]),
        )
        for (r_code, op_no), items in route_items:
            op_name, op_content = op_map.get(f"{r_code}:{op_no}", (f"工序{op_no}", ""))
            final_items = items or fallback_items(r_code, op_no, op_name, op_content)
            for idx, item in enumerate(final_items, start=1):
                tool_value = item.get("tool", "")
                if not tool_value and item.get("tool_name"):
                    tool_value = tool_map.get(item["tool_name"], item["tool_name"])
                code = f"CI-{route_code.replace('RtDJ', 'DJ')}-{op_no}-{idx:02d}"
                ws_items.append(
                    [
                        plan_code,
                        route_code,
                        route["route_name"],
                        op_no,
                        op_name,
                        idx,
                        code,
                        item["name"],
                        tool_value,
                        item["feature"],
                        item["std"],
                        item["upper"],
                        item["lower"],
                        item["method"],
                        item["note"] or op_content,
                    ]
                )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    wb.close()
    return output_path


def parse_args():
    parser = argparse.ArgumentParser(description="生成大疆无人机工艺路线-工序检查项配置工作簿")
    parser.add_argument("--route-workbook", type=Path, required=True, help="最新工艺路线工作簿")
    parser.add_argument("--resource-workbook", type=Path, required=True, help="工厂资源工作簿")
    parser.add_argument("--output", type=Path, required=True, help="输出xlsx路径")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    output_path = build_workbook(args.route_workbook.resolve(), args.resource_workbook.resolve(), args.output.resolve())
    print(f"[OK] 已生成: {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
