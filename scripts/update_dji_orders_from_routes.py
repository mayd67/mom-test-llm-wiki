from __future__ import annotations

import argparse
import sys
from pathlib import Path

from openpyxl import load_workbook

try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass


TARGET_SHEETS = ("生产订单", "备料清单")


def normalize(value) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    return str(value).strip()


def read_table(worksheet) -> tuple[list[str], list[dict[str, str]]]:
    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = [normalize(cell) for cell in rows[0]]
    result: list[dict[str, str]] = []
    for raw_row in rows[1:]:
        item: dict[str, str] = {}
        has_value = False
        for idx, header in enumerate(headers):
            if not header:
                continue
            value = normalize(raw_row[idx] if idx < len(raw_row) else "")
            if value:
                has_value = True
            item[header] = value
        if has_value:
            result.append(item)
    return headers, result


def snapshot_rows(worksheet) -> list[tuple[str, ...]]:
    result: list[tuple[str, ...]] = []
    for row in worksheet.iter_rows(values_only=True):
        values = tuple(normalize(cell) for cell in row)
        if any(values):
            result.append(values)
    return result


def clear_and_write_rows(worksheet, rows: list[list[str]]) -> None:
    max_col = max(worksheet.max_column, max((len(row) for row in rows), default=0))
    for r in range(2, worksheet.max_row + 1):
        for c in range(1, max_col + 1):
            worksheet.cell(r, c).value = None
    for r_idx, row in enumerate(rows, start=2):
        for c_idx, value in enumerate(row, start=1):
            worksheet.cell(r_idx, c_idx).value = value


def build_route_maps(route_wb):
    _, route_rows = read_table(route_wb["工艺路线"])
    _, route_op_rows = read_table(route_wb["工艺路线工序"])
    _, route_material_rows = read_table(route_wb["工艺路线工序物料"])

    route_by_code: dict[str, dict[str, str]] = {}
    route_by_material: dict[str, dict[str, str]] = {}
    for row in route_rows:
        code = normalize(row.get("*编码"))
        material_code = normalize(row.get("物料编码"))
        if code:
            route_by_code[code] = row
        if material_code:
            route_by_material[material_code] = row

    op_name_by_route_op: dict[tuple[str, str], str] = {}
    for row in route_op_rows:
        key = (normalize(row.get("*工艺路线编码")), normalize(row.get("*工序号")))
        op_name_by_route_op[key] = normalize(row.get("*工序名称"))

    materials_by_route: dict[str, list[dict[str, str]]] = {}
    for row in route_material_rows:
        route_code = normalize(row.get("*工艺路线编码"))
        if route_code:
            materials_by_route.setdefault(route_code, []).append(row)

    return route_by_code, route_by_material, op_name_by_route_op, materials_by_route


def row_to_header_order(row: dict[str, str], headers: list[str]) -> list[str]:
    return [normalize(row.get(header, "")) for header in headers]


def to_number_text(value: str) -> str:
    try:
        num = float(value)
    except ValueError:
        return value
    if num.is_integer():
        return str(int(num))
    return str(round(num, 3)).rstrip("0").rstrip(".")


def build_pick_rows(
    order_rows: list[dict[str, str]],
    op_name_by_route_op: dict[tuple[str, str], str],
    materials_by_route: dict[str, list[dict[str, str]]],
) -> list[dict[str, str]]:
    result: list[dict[str, str]] = []
    for order_row in order_rows:
        order_code = normalize(order_row.get("*编码"))
        route_code = normalize(order_row.get("工艺路线编码"))
        route_version = normalize(order_row.get("工艺路线版本号"))
        plan_qty_text = normalize(order_row.get("*计划数量"))
        try:
            plan_qty = float(plan_qty_text)
        except ValueError:
            plan_qty = 0.0

        for material_row in materials_by_route.get(route_code, []):
            unit_qty_text = normalize(material_row.get("数量"))
            try:
                unit_qty = float(unit_qty_text)
            except ValueError:
                unit_qty = 0.0
            demand_qty = plan_qty * unit_qty
            op_no = normalize(material_row.get("*工序号"))
            result.append(
                {
                    "*生产订单编码": order_code,
                    "*物料版本号": normalize(material_row.get("*物料版本号")) or route_version,
                    "*物料编码": normalize(material_row.get("*物料编码")),
                    "工序名称": op_name_by_route_op.get((route_code, op_no), ""),
                    "*需求数量": to_number_text(str(demand_qty)),
                    "*子件比例": to_number_text(unit_qty_text),
                    "替换件物料版本号": "",
                    "替换件物料编码": "",
                    "是否必须装入": "是",
                    "工序编码": op_no,
                }
            )
    return result


def build_updated_orders(order_rows: list[dict[str, str]], route_by_code, route_by_material) -> list[dict[str, str]]:
    updated: list[dict[str, str]] = []
    for row in order_rows:
        material_code = normalize(row.get("*物料编码"))
        current_route_code = normalize(row.get("工艺路线编码"))
        route_meta = route_by_code.get(current_route_code) or route_by_material.get(material_code)
        new_row = dict(row)
        if route_meta:
            new_row["工艺路线版本号"] = normalize(route_meta.get("*版本号")) or new_row.get("工艺路线版本号", "")
            new_row["工艺路线编码"] = normalize(route_meta.get("*编码")) or new_row.get("工艺路线编码", "")
            if normalize(route_meta.get("*工厂组织")):
                new_row["*所属组织"] = normalize(route_meta.get("*工厂组织"))
        updated.append(new_row)
    return updated


def fallback_output_path(output_path: Path) -> Path:
    return output_path.with_name(f"{output_path.stem}_按最新工艺路线更新{output_path.suffix}")


def parse_args():
    parser = argparse.ArgumentParser(description="根据最新工艺路线更新生产订单与备料清单")
    parser.add_argument("--orders", type=Path, required=True, help="生产订单模板工作簿")
    parser.add_argument("--routes", type=Path, required=True, help="最新工艺路线工作簿")
    parser.add_argument("--output", type=Path, default=None, help="输出路径，默认覆盖原文件")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    order_path = args.orders.resolve()
    output_path = (args.output or args.orders).resolve()
    route_path = args.routes.resolve()

    order_wb = load_workbook(order_path)
    route_wb = load_workbook(route_path, data_only=True)

    try:
        original_non_target = {
            name: snapshot_rows(order_wb[name])
            for name in order_wb.sheetnames
            if name not in TARGET_SHEETS
        }

        order_headers, order_rows = read_table(order_wb["生产订单"])
        pick_headers, _pick_rows = read_table(order_wb["备料清单"])
        route_by_code, route_by_material, op_name_by_route_op, materials_by_route = build_route_maps(route_wb)

        updated_order_rows = build_updated_orders(order_rows, route_by_code, route_by_material)
        updated_pick_rows = build_pick_rows(updated_order_rows, op_name_by_route_op, materials_by_route)

        clear_and_write_rows(
            order_wb["生产订单"],
            [row_to_header_order(row, order_headers) for row in updated_order_rows],
        )
        clear_and_write_rows(
            order_wb["备料清单"],
            [row_to_header_order(row, pick_headers) for row in updated_pick_rows],
        )

        actual_output = output_path
        try:
            order_wb.save(actual_output)
        except PermissionError:
            actual_output = fallback_output_path(output_path)
            order_wb.save(actual_output)

        verify_wb = load_workbook(actual_output, data_only=True)
        try:
            for name, rows in original_non_target.items():
                if snapshot_rows(verify_wb[name]) != rows:
                    raise RuntimeError(f"非目标sheet发生变化: {name}")

            actual_orders = snapshot_rows(verify_wb["生产订单"])[1:]
            expected_orders = [
                tuple(row_to_header_order(row, order_headers))
                for row in updated_order_rows
            ]
            if actual_orders != expected_orders:
                raise RuntimeError("生产订单sheet写入校验失败")

            actual_picks = snapshot_rows(verify_wb["备料清单"])[1:]
            expected_picks = [
                tuple(row_to_header_order(row, pick_headers))
                for row in updated_pick_rows
            ]
            if actual_picks != expected_picks:
                raise RuntimeError("备料清单sheet写入校验失败")
        finally:
            verify_wb.close()

        print(f"[OK] 已生成: {actual_output}")
        print(f"[OK] 生产订单: {len(updated_order_rows)} rows")
        print(f"[OK] 备料清单: {len(updated_pick_rows)} rows")
        return 0
    finally:
        order_wb.close()
        route_wb.close()


if __name__ == "__main__":
    raise SystemExit(main())
