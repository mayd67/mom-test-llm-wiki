from __future__ import annotations

import argparse
import json
import re
import shutil
import subprocess
import sys
import time
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from datetime import datetime
from html.parser import HTMLParser
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
RAW_DIR = ROOT / "raw"
WIKI_DIR = ROOT / "wiki"
OUTPUTS_DIR = ROOT / "outputs"
PLAN_PATH = OUTPUTS_DIR / "raw_to_wiki_ingest_plan.json"
STATE_PATH = OUTPUTS_DIR / "raw_to_wiki_ingest_state.json"
AUTO_MARKER = "<!-- raw-to-wiki-ingest:auto -->"

PROFILE_LEGACY = "legacy"
PROFILE_TEAM_TEMPLATE = "team-template"
PROFILE_MIXED = "mixed"
PROFILE_UNKNOWN = "unknown"

TEXT_EXTENSIONS = {".md", ".txt", ".yaml", ".yml", ".html", ".htm"}
ATTACHMENT_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".mp4"}
SKIP_NAMES = {".gitkeep", ".gitignore"}
SKIP_PARTS = {".git", ".hg", ".svn", ".trae", "__pycache__"}
LEGACY_WIKI_RAW_DIR = "99_原始资料"
LEGACY_WIKI_MIRROR_DIR = "99_原始资料镜像"
LINK_RE = re.compile(r"(!?\[[^\]]*\]\()([^)]+)(\))")

LEGACY_MIRROR_ROOTS = {
    "需求文档": WIKI_DIR / "01_MOM产品相关文档" / "01_产品需求文档" / "99_原始资料镜像",
    "产品资料": WIKI_DIR / "01_MOM产品相关文档" / "02_数据模型" / "10_自动入库镜像",
    "接口资料": WIKI_DIR / "01_MOM产品相关文档" / "02_数据模型" / "10_自动入库镜像" / "接口资料",
    "测试资料": WIKI_DIR / "01_MOM产品相关文档" / "04_测试用例" / "99_原始资料镜像",
    "测试流程规范": WIKI_DIR / "01_MOM产品相关文档" / "00_方法规范" / "99_原始资料镜像" / "测试流程规范",
    "截图附件": WIKI_DIR / "01_MOM产品相关文档" / "00_方法规范" / "99_原始资料镜像" / "截图附件",
}
LEGACY_PROJECT_DOC_ROOTS = {
    "BJHX项目资料": WIKI_DIR / "02_HX项目资料" / "99_原始资料镜像",
    "QD项目资料": WIKI_DIR / "03_QD项目资料" / "99_原始资料",
}
DEFAULT_LEGACY_PROJECT_DOC_ROOT = WIKI_DIR / "02_HX项目资料" / "99_原始资料镜像"
SUPPORTED_RAW_GROUPS = set(LEGACY_MIRROR_ROOTS) | {"项目文档"}

TEAM_TEMPLATE_BUILDERS = [
    ("需求文档", ROOT / "scripts" / "build_requirements_wiki.py"),
    ("raw-markdown", ROOT / "scripts" / "build_raw_markdown_mirror.py"),
    ("raw-office", ROOT / "scripts" / "build_raw_office_mirror.py"),
    ("raw-attachments", ROOT / "scripts" / "build_raw_attachment_catalog.py"),
    ("kmmom-product-docs", ROOT / "scripts" / "build_kmmom_product_docs_wiki.py"),
    ("kmmom-data-dictionary", ROOT / "scripts" / "build_kmmom_data_dictionary.py"),
    ("test-case-wiki", ROOT / "scripts" / "build_test_case_wiki.py"),
    ("coverage-matrix", ROOT / "scripts" / "build_wiki_coverage_matrix.py"),
    ("sync-meta", ROOT / "scripts" / "sync_wiki_meta.py"),
]


@dataclass
class Route:
    raw_path: str
    group: str
    action: str
    handler: str
    target_path: str | None = None
    note: str | None = None


@dataclass(frozen=True)
class SourceContext:
    kind: str
    root: Path
    group: str


def now_iso() -> str:
    return datetime.now().replace(microsecond=0).isoformat()


def relative_posix(path: Path, base: Path = ROOT) -> str:
    return path.relative_to(base).as_posix()


def read_text_with_fallback(path: Path) -> str:
    for encoding in ("utf-8", "utf-8-sig", "gb18030"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_text(encoding="utf-8", errors="replace")


class TextOnlyHTMLParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        if data.strip():
            self.parts.append(data.strip())

    def text(self) -> str:
        return "\n".join(self.parts)


def detect_profile() -> str:
    if not WIKI_DIR.exists():
        return PROFILE_UNKNOWN
    top_dirs = {path.name for path in WIKI_DIR.iterdir() if path.is_dir()}
    has_legacy = {"01_MOM产品相关文档", "02_HX项目资料"} <= top_dirs
    has_team_template = {"01_通用规范", "02_测试标准&模板", "03_业务系统"} & top_dirs
    if has_legacy and has_team_template:
        return PROFILE_MIXED
    if has_legacy:
        return PROFILE_LEGACY
    if has_team_template:
        return PROFILE_TEAM_TEMPLATE
    return PROFILE_UNKNOWN


def resolve_profile(requested: str) -> str:
    detected = detect_profile()
    if requested != "auto":
        return requested
    return detected


def legacy_wiki_source_root(path: Path) -> Path | None:
    try:
        rel = path.relative_to(WIKI_DIR.resolve())
    except ValueError:
        return None
    if len(rel.parts) < 2 or rel.parts[1] != LEGACY_WIKI_RAW_DIR:
        return None
    return (WIKI_DIR / rel.parts[0] / LEGACY_WIKI_RAW_DIR).resolve()


def source_context(path: Path) -> SourceContext | None:
    resolved = path.resolve()
    try:
        rel = resolved.relative_to(RAW_DIR.resolve())
    except ValueError:
        legacy_root = legacy_wiki_source_root(resolved)
        if legacy_root is None:
            return None
        return SourceContext(kind="legacy-wiki", root=legacy_root, group="项目文档")

    if not rel.parts:
        return None
    group = rel.parts[0]
    if group not in SUPPORTED_RAW_GROUPS:
        return None
    return SourceContext(kind="raw", root=RAW_DIR.resolve(), group=group)


def normalize_input_path(value: str) -> Path:
    path = Path(value)
    if not path.is_absolute():
        path = (ROOT / path).resolve()
    else:
        path = path.resolve()
    if source_context(path) is None:
        raise SystemExit(f"Input path must stay under raw/ or wiki/*/{LEGACY_WIKI_RAW_DIR}/: {value}")
    return path


def group_relative_path(path: Path) -> Path:
    context = source_context(path)
    if context is None:
        raise SystemExit(f"Unsupported source path: {relative_posix(path)}")
    if context.kind == "legacy-wiki":
        return path.relative_to(context.root)
    return path.relative_to(RAW_DIR / context.group)


def iter_selected_files(inputs: list[str] | None) -> list[Path]:
    if not inputs:
        return sorted(path for path in RAW_DIR.rglob("*") if path.is_file())
    result: list[Path] = []
    for value in inputs:
        path = normalize_input_path(value)
        if path.is_file():
            result.append(path)
            continue
        if path.is_dir():
            result.extend(sorted(item for item in path.rglob("*") if item.is_file()))
            continue
        raise SystemExit(f"Path does not exist: {value}")
    return sorted(dict.fromkeys(result))


def should_skip(path: Path) -> bool:
    if path.name in SKIP_NAMES or path.name.startswith(".") or path.name.startswith("~$"):
        return True
    if any(part in SKIP_PARTS for part in path.parts):
        return True
    return source_context(path) is None


def top_group(path: Path) -> str:
    context = source_context(path)
    if context is None:
        raise SystemExit(f"Unsupported source path: {relative_posix(path)}")
    return context.group


def legacy_mirror_root(path: Path) -> Path:
    context = source_context(path)
    if context is None:
        raise SystemExit(f"Unsupported source path: {relative_posix(path)}")
    if context.kind == "legacy-wiki":
        return context.root.parent / LEGACY_WIKI_MIRROR_DIR

    group = context.group
    if group == "项目文档":
        rel_parts = group_relative_path(path).parts
        project_folder = rel_parts[0] if rel_parts else ""
        return LEGACY_PROJECT_DOC_ROOTS.get(project_folder, DEFAULT_LEGACY_PROJECT_DOC_ROOT)

    root = LEGACY_MIRROR_ROOTS.get(group)
    if root is None:
        raise SystemExit(f"Unsupported raw group in legacy profile: {group}")
    return root


def legacy_mirror_path(path: Path) -> Path:
    rel = group_relative_path(path)
    target_root = legacy_mirror_root(path)
    return target_root.joinpath(*rel.parts).with_suffix(".md")


def legacy_attachment_index_path(path: Path) -> Path:
    rel_dir = group_relative_path(path).parent
    target_root = legacy_mirror_root(path)
    return target_root.joinpath(*rel_dir.parts) / "附件索引.md"


def route_for_legacy(path: Path) -> Route:
    context = source_context(path)
    if context is None:
        raise SystemExit(f"Unsupported source path: {relative_posix(path)}")
    group = top_group(path)
    rel = relative_posix(path)
    if context.kind == "raw" and group == "项目文档":
        rel_parts = group_relative_path(path).parts
        if rel_parts and rel_parts[0] == "QD项目资料":
            return Route(
                raw_path=rel,
                group=group,
                action="copy_raw_file",
                handler="legacy_copy_raw_file",
                target_path=relative_posix(legacy_mirror_root(path).joinpath(*rel_parts)),
                note="QD project docs are copied from raw into wiki as original files.",
            )
    extension = path.suffix.lower()
    if extension in ATTACHMENT_EXTENSIONS:
        return Route(
            raw_path=rel,
            group=group,
            action="index_attachments",
            handler="legacy_attachment_index",
            target_path=relative_posix(legacy_attachment_index_path(path)),
            note="Binary attachments stay in raw and are exposed through a wiki index page.",
        )
    return Route(
        raw_path=rel,
        group=group,
        action="mirror_file",
        handler="legacy_mirror_file",
        target_path=relative_posix(legacy_mirror_path(path)),
        note="Safe legacy mode writes only to dedicated auto-ingest mirror folders.",
    )


def builders_for_team_template(files: list[Path]) -> list[Path]:
    selected: list[Path] = []
    all_extensions = {path.suffix.lower() for path in files}
    raw_groups = {top_group(path) for path in files}
    file_strings = {relative_posix(path) for path in files}

    def add(name: str) -> None:
        for key, script in TEAM_TEMPLATE_BUILDERS:
            if key == name and script not in selected:
                selected.append(script)

    if "需求文档" in raw_groups:
        add("需求文档")
    if all_extensions & TEXT_EXTENSIONS:
        add("raw-markdown")
    if all_extensions & {".docx", ".xlsx"}:
        add("raw-office")
    if all_extensions & ATTACHMENT_EXTENSIONS:
        add("raw-attachments")
    if any(item.startswith("raw/产品资料/km-mom-docs/") for item in file_strings):
        add("kmmom-product-docs")
    if any("raw/产品资料/km-mom-docs/03-development/datamodel-design/" in item for item in file_strings):
        add("kmmom-data-dictionary")
    if any(item.startswith("raw/测试资料/01_测试用例/") for item in file_strings):
        add("test-case-wiki")
    if selected:
        add("coverage-matrix")
        add("sync-meta")
    return selected


def plan_for_team_template(files: list[Path]) -> list[Route]:
    routes: list[Route] = []
    builders = [relative_posix(path) for path in builders_for_team_template(files)]
    note = "Profile uses the new team-template builders; apply mode runs the required batch scripts."
    for path in files:
        routes.append(
            Route(
                raw_path=relative_posix(path),
                group=top_group(path),
                action="run_builders",
                handler="team_template_builders",
                target_path=None,
                note=f"{note} Builders: {', '.join(builders)}" if builders else note,
            )
        )
    return routes


def build_plan(profile: str, files: list[Path]) -> dict:
    if profile == PROFILE_LEGACY:
        routes = [route_for_legacy(path) for path in files if not should_skip(path)]
    elif profile == PROFILE_TEAM_TEMPLATE:
        routes = plan_for_team_template([path for path in files if not should_skip(path)])
    elif profile == PROFILE_MIXED:
        routes = []
    else:
        routes = []

    summary = Counter(route.action for route in routes)
    handler_summary = Counter(route.handler for route in routes)
    plan = {
        "generated_at": now_iso(),
        "requested_profile": profile,
        "detected_profile": detect_profile(),
        "raw_file_count": len(files),
        "route_count": len(routes),
        "summary": dict(summary),
        "handlers": dict(handler_summary),
        "routes": [asdict(route) for route in routes],
    }
    return plan


def write_json(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def load_json(path: Path) -> dict | None:
    if not path.exists():
        return None
    return json.loads(path.read_text(encoding="utf-8"))


def build_snapshot(files: list[Path]) -> dict[str, dict[str, int]]:
    snapshot: dict[str, dict[str, int]] = {}
    for path in files:
        if should_skip(path):
            continue
        stat = path.stat()
        snapshot[relative_posix(path)] = {"mtime_ns": stat.st_mtime_ns, "size": stat.st_size}
    return snapshot


def diff_snapshot(old: dict[str, dict[str, int]] | None, new: dict[str, dict[str, int]]) -> tuple[list[str], list[str], list[str]]:
    if old is None:
        return sorted(new), [], []
    old_keys = set(old)
    new_keys = set(new)
    added = sorted(new_keys - old_keys)
    removed = sorted(old_keys - new_keys)
    modified = sorted(
        key
        for key in old_keys & new_keys
        if old[key]["mtime_ns"] != new[key]["mtime_ns"] or old[key]["size"] != new[key]["size"]
    )
    return added, removed, modified


def shift_markdown_headings(text: str) -> str:
    lines = []
    for line in text.splitlines():
        if line.startswith("#"):
            lines.append("#" + line)
        else:
            lines.append(line)
    return "\n".join(lines).strip()


def title_for_raw_file(path: Path) -> str:
    return f"原始资料-{top_group(path)}-{path.stem}"


def yaml_list(lines: list[str], indent: int = 2) -> list[str]:
    prefix = " " * indent
    return [f"{prefix}- {line}" for line in lines]


def relative_markdown_link(from_path: Path, to_path: Path) -> str:
    origin = from_path.parent.relative_to(ROOT)
    rel = Path(*([".."] * len(origin.parts))) / to_path.relative_to(ROOT) if origin.parts else to_path.relative_to(ROOT)
    return rel.as_posix()


def rewrite_markdown_links(text: str, raw_file: Path, target: Path) -> str:
    def replace(match: re.Match[str]) -> str:
        prefix, link_target, suffix = match.groups()
        if link_target.startswith(("http://", "https://", "mailto:", "#")):
            return match.group(0)

        wrapped = link_target.startswith("<") and link_target.endswith(">")
        clean_target = link_target[1:-1] if wrapped else link_target
        resolved = (raw_file.parent / clean_target).resolve()
        if not resolved.exists():
            return match.group(0)

        rewritten = relative_markdown_link(target, resolved)
        if wrapped:
            rewritten = f"<{rewritten}>"
        return f"{prefix}{rewritten}{suffix}"

    return LINK_RE.sub(replace, text)


def render_frontmatter(title: str, summary: str, sources: list[str], tags: list[str]) -> list[str]:
    lines = ["---", f"title: {title}", "type: manual", "status: active", "tags:"]
    lines.extend(yaml_list(tags))
    lines.append(f"summary: {summary}")
    lines.append("source:")
    lines.extend(yaml_list(sources))
    lines.append(f"updated: {datetime.now().date().isoformat()}")
    lines.append("---")
    return lines


def extract_html(path: Path) -> str:
    parser = TextOnlyHTMLParser()
    parser.feed(read_text_with_fallback(path))
    return parser.text()


def extract_docx(path: Path) -> str:
    try:
        from docx import Document  # type: ignore
    except ImportError:
        return "未提取正文：当前环境缺少 python-docx，可保留原文并后续补装依赖。"

    doc = Document(path)
    lines: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            lines.append(text)
    for table_index, table in enumerate(doc.tables, start=1):
        rows = []
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " / ") for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            lines.append("")
            lines.append(f"## 表格 {table_index}")
            lines.extend(rows)
    return "\n".join(lines).strip()


def extract_xlsx(path: Path) -> str:
    try:
        from openpyxl import load_workbook  # type: ignore
    except ImportError:
        return "未提取正文：当前环境缺少 openpyxl，可保留原文并后续补装依赖。"

    workbook = load_workbook(path, read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in workbook.worksheets:
        lines.append(f"## 工作表：{sheet.title}")
        row_count = 0
        for row in sheet.iter_rows(values_only=True):
            values = [str(value).strip() for value in row if value not in (None, "")]
            if not values:
                continue
            row_count += 1
            if row_count > 40:
                lines.append("- ...")
                break
            lines.append(f"- {' | '.join(values[:12])}")
        if row_count == 0:
            lines.append("- 空表")
        lines.append("")
    return "\n".join(lines).strip()


def extract_body(path: Path) -> tuple[str, str]:
    extension = path.suffix.lower()
    if extension == ".md":
        return shift_markdown_headings(read_text_with_fallback(path)), "markdown"
    if extension in {".txt", ".yaml", ".yml"}:
        return read_text_with_fallback(path).strip(), "text"
    if extension in {".html", ".htm"}:
        return extract_html(path), "html"
    if extension == ".docx":
        return extract_docx(path), "docx"
    if extension == ".xlsx":
        return extract_xlsx(path), "xlsx"
    return "当前文件类型暂不支持正文提取，已保留来源路径供后续人工整理。", "unsupported"


def render_legacy_mirror(path: Path, target: Path) -> str:
    body, body_kind = extract_body(path)
    if body_kind == "markdown":
        body = rewrite_markdown_links(body, path, target)
    raw_rel = relative_posix(path)
    title = title_for_raw_file(path)
    summary = f"自动镜像 {raw_rel}，作为当前 legacy wiki 结构下的安全入库入口。"
    lines = render_frontmatter(title, summary, [raw_rel], ["testing", "raw", "mirror", "auto-ingest"])
    lines.extend(
        [
            "",
            AUTO_MARKER,
            "",
            f"# {title}",
            "",
            "## 来源说明",
            "",
            f"- 原始路径：`{raw_rel}`",
            f"- 目标路径：`{relative_posix(target)}`",
            f"- 提取方式：`{body_kind}`",
            "- 说明：本页由统一入库脚本自动生成，不覆盖人工整理页。",
            "",
            "## 原始内容镜像",
            "",
            body or "未提取到可展示正文。",
            "",
        ]
    )
    return "\n".join(lines)


def render_attachment_index(raw_dir: Path, target: Path) -> str:
    group = top_group(raw_dir)
    raw_rel_dir = relative_posix(raw_dir)
    title = f"附件索引-{group}-{raw_dir.name}"
    tags = ["testing", "raw", "attachment", "auto-ingest"]
    lines = render_frontmatter(title, f"汇总 {raw_rel_dir} 下的附件入口。", [f"{raw_rel_dir}/"], tags)
    lines.extend(
        [
            "",
            AUTO_MARKER,
            "",
            f"# {title}",
            "",
            "## 来源说明",
            "",
            f"- 原始目录：`{raw_rel_dir}`",
            f"- 目标路径：`{relative_posix(target)}`",
            "- 说明：附件原件继续保留在来源目录，仅在 wiki 中建立索引入口。",
            "",
            "## 附件列表",
            "",
        ]
    )
    attachments = sorted(
        item for item in raw_dir.iterdir() if item.is_file() and item.suffix.lower() in ATTACHMENT_EXTENSIONS
    )
    if not attachments:
        lines.append("- 当前目录下没有附件。")
    for item in attachments:
        rel_link = relative_markdown_link(target, item)
        lines.append(f"- [{item.name}]({rel_link})")
    lines.append("")
    return "\n".join(lines)


def write_auto_page(path: Path, content: str) -> str:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        old = path.read_text(encoding="utf-8", errors="replace")
        if AUTO_MARKER not in old:
            return "skipped-manual"
    path.write_text(content.rstrip() + "\n", encoding="utf-8")
    return "written"


def copy_raw_file(source: Path, target: Path) -> str:
    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, target)
    return "copied"


def apply_legacy(routes: list[Route]) -> dict[str, int]:
    result = Counter()
    attachment_targets: dict[Path, list[Path]] = defaultdict(list)

    for route in routes:
        raw_path = ROOT / route.raw_path
        if route.action == "copy_raw_file" and route.target_path:
            target = ROOT / route.target_path
            status = copy_raw_file(raw_path, target)
            result[status] += 1
            continue
        if route.action == "index_attachments" and route.target_path:
            attachment_targets[ROOT / route.target_path].append(raw_path)
            continue
        if route.action != "mirror_file" or not route.target_path:
            result["skipped"] += 1
            continue
        target = ROOT / route.target_path
        status = write_auto_page(target, render_legacy_mirror(raw_path, target))
        result[status] += 1

    for target, raw_files in attachment_targets.items():
        raw_dir = raw_files[0].parent
        status = write_auto_page(target, render_attachment_index(raw_dir, target))
        result[status] += 1

    return dict(result)


def apply_team_template(files: list[Path]) -> dict[str, int]:
    builders = builders_for_team_template(files)
    result = Counter()
    for builder in builders:
        subprocess.run([sys.executable, str(builder)], cwd=ROOT, check=True)
        result["builders_run"] += 1
    return dict(result)


def print_plan(plan: dict) -> None:
    print(f"profile={plan['requested_profile']} detected={plan['detected_profile']}")
    print(f"raw_file_count={plan['raw_file_count']} route_count={plan['route_count']}")
    if plan["summary"]:
        for action, count in sorted(plan["summary"].items()):
            print(f"- {action}: {count}")
    if plan["handlers"]:
        for handler, count in sorted(plan["handlers"].items()):
            print(f"  handler {handler}: {count}")


def select_files_for_apply(args: argparse.Namespace) -> tuple[list[Path], dict | None, dict | None]:
    files = iter_selected_files(args.paths)
    if not args.incremental:
        return files, None, None

    snapshot = build_snapshot(files)
    old_snapshot = load_json(Path(args.state_file))
    added, removed, modified = diff_snapshot(old_snapshot, snapshot)
    changed = added + modified
    selected = [ROOT / item for item in changed]
    if removed:
        print(f"removed_raw_files={len(removed)} (wiki mirrors are not auto-deleted)")
    return selected, snapshot, old_snapshot


def command_plan(args: argparse.Namespace) -> int:
    profile = resolve_profile(args.profile)
    if profile in {PROFILE_UNKNOWN, PROFILE_MIXED}:
        raise SystemExit(f"Cannot build plan automatically for profile: {profile}")
    files = iter_selected_files(args.paths)
    plan = build_plan(profile, files)
    write_json(Path(args.plan_file), plan)
    print_plan(plan)
    return 0


def command_apply(args: argparse.Namespace) -> int:
    profile = resolve_profile(args.profile)
    if profile in {PROFILE_UNKNOWN, PROFILE_MIXED}:
        raise SystemExit(
            f"Cannot apply ingestion for profile {profile}. Use --profile to force a mode after manual review."
        )

    files, snapshot, _old_snapshot = select_files_for_apply(args)
    if not files:
        print("No changed raw files detected.")
        if snapshot is not None:
            write_json(Path(args.state_file), snapshot)
        return 0

    plan = build_plan(profile, files)
    write_json(Path(args.plan_file), plan)
    print_plan(plan)

    if args.dry_run:
        return 0

    if profile == PROFILE_LEGACY:
        result = apply_legacy([Route(**route) for route in plan["routes"]])
        print(json.dumps(result, ensure_ascii=False, indent=2))
    else:
        result = apply_team_template(files)
        print(json.dumps(result, ensure_ascii=False, indent=2))

    if snapshot is not None:
        write_json(Path(args.state_file), snapshot)
    return 0


def command_watch(args: argparse.Namespace) -> int:
    args.incremental = True
    while True:
        try:
            command_apply(args)
            time.sleep(args.interval)
        except KeyboardInterrupt:
            print("Watch stopped.")
            return 0


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Profile-aware raw -> wiki ingestion entrypoint.")
    parser.add_argument("--profile", choices=["auto", PROFILE_LEGACY, PROFILE_TEAM_TEMPLATE], default="auto")
    parser.add_argument("--plan-file", default=str(PLAN_PATH))
    parser.add_argument("--state-file", default=str(STATE_PATH))

    subparsers = parser.add_subparsers(dest="command", required=True)

    plan_parser = subparsers.add_parser("plan", help="Only compute the ingestion plan.")
    plan_parser.add_argument("paths", nargs="*")
    plan_parser.set_defaults(func=command_plan)

    apply_parser = subparsers.add_parser("apply", help="Apply ingestion for the selected raw files.")
    apply_parser.add_argument("paths", nargs="*")
    apply_parser.add_argument("--incremental", action="store_true")
    apply_parser.add_argument("--dry-run", action="store_true")
    apply_parser.set_defaults(func=command_apply)

    watch_parser = subparsers.add_parser("watch", help="Poll raw/ and ingest incremental changes.")
    watch_parser.add_argument("paths", nargs="*")
    watch_parser.add_argument("--dry-run", action="store_true")
    watch_parser.add_argument("--interval", type=int, default=30)
    watch_parser.set_defaults(func=command_watch)

    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()
    return args.func(args)


if __name__ == "__main__":
    raise SystemExit(main())
