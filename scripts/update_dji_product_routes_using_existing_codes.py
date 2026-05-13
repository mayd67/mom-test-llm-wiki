from __future__ import annotations

import argparse
import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook

try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass


SOURCE_ROUTE_SHEETS = ("无人机总装工艺路线", "电控印制板工艺路线")
ROUTE_SHEETS_TO_REWRITE = (
    "工艺路线工序",
    "工艺路线工序序列",
    "工艺路线工序物料",
    "工艺路线工步",
)


def normalize(value) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    return str(value).strip()


def read_table(worksheet) -> list[dict[str, str]]:
    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [normalize(cell) for cell in rows[0]]
    result: list[dict[str, str]] = []
    for raw_row in rows[1:]:
        item: dict[str, str] = {}
        has_value = False
        for idx, header in enumerate(headers):
            if not header:
                continue
            value = normalize(raw_row[idx] if idx < len(raw_row) else "")
            if value:
                has_value = True
            item[header] = value
        if has_value:
            result.append(item)
    return result


def op_type(name: str) -> str:
    if any(keyword in name for keyword in ("检验", "测试", "校对")):
        return "检验"
    return "加工"


def route_content(op_row: dict[str, str]) -> str:
    parts = []
    involved = normalize(op_row.get("涉及物料"))
    if involved:
        parts.append(f"涉及物料：{involved}")
    control = normalize(op_row.get("关键控制点"))
    if control:
        parts.append(f"关键控制点：{control}")
    record = normalize(op_row.get("质量记录/防错")) or normalize(op_row.get("检测/记录要求"))
    if record:
        parts.append(f"记录要求：{record}")
    return "；".join(parts)


def snapshot_rows(worksheet) -> list[tuple[str, ...]]:
    data: list[tuple[str, ...]] = []
    for row in worksheet.iter_rows(values_only=True):
        values = tuple(normalize(cell) for cell in row)
        if any(values):
            data.append(values)
    return data


def clear_and_write_rows(worksheet, rows: list[list[str]]) -> None:
    max_col = max(worksheet.max_column, max((len(row) for row in rows), default=0))
    for r in range(2, worksheet.max_row + 1):
        for c in range(1, max_col + 1):
            worksheet.cell(r, c).value = None
    for r_idx, row in enumerate(rows, start=2):
        for c_idx, value in enumerate(row, start=1):
            worksheet.cell(r_idx, c_idx).value = value


def build_material_code_map(product_wb) -> dict[str, str]:
    sheet = product_wb["物料"]
    result: dict[str, str] = {}
    for row in read_table(sheet):
        drawing = normalize(row.get("图号"))
        code = normalize(row.get("*编码"))
        name = normalize(row.get("*名称"))
        if drawing and code:
            result[drawing] = code
        if name and code and name not in result:
            result[name] = code
    return result


def build_route_rows_by_name(product_wb) -> dict[str, dict[str, str]]:
    sheet = product_wb["工艺路线"]
    result: dict[str, dict[str, str]] = {}
    for row in read_table(sheet):
        result[normalize(row.get("*名称"))] = row
    return result


def build_existing_exec_op_map(product_wb) -> dict[tuple[str, str], dict[str, str]]:
    route_sheet = product_wb["工艺路线"]
    op_sheet = product_wb["工艺路线工序"]

    route_name_by_code = {
        normalize(row.get("*编码")): normalize(row.get("*名称"))
        for row in read_table(route_sheet)
    }

    result: dict[tuple[str, str], dict[str, str]] = {}
    for row in read_table(op_sheet):
        route_code = normalize(row.get("*工艺路线编码"))
        route_name = route_name_by_code.get(route_code, route_code)
        op_name = normalize(row.get("*工序名称"))
        if route_name and op_name:
            result[(route_name, op_name)] = row
    return result


def build_old_op_materials_map(product_wb) -> dict[tuple[str, str], list[dict[str, str]]]:
    route_sheet = product_wb["工艺路线"]
    op_sheet = product_wb["工艺路线工序"]
    mat_sheet = product_wb["工艺路线工序物料"]

    route_name_by_code = {
        normalize(row.get("*编码")): normalize(row.get("*名称"))
        for row in read_table(route_sheet)
    }

    exec_op_key_by_route_opno: dict[tuple[str, str], tuple[str, str]] = {}
    for row in read_table(op_sheet):
        route_code = normalize(row.get("*工艺路线编码"))
        route_name = route_name_by_code.get(route_code, route_code)
        op_no = normalize(row.get("*工序号"))
        op_name = normalize(row.get("*工序名称"))
        exec_op_key_by_route_opno[(route_code, op_no)] = (route_name, op_name)

    grouped: dict[tuple[str, str], list[dict[str, str]]] = {}
    for row in read_table(mat_sheet):
        route_code = normalize(row.get("*工艺路线编码"))
        op_no = normalize(row.get("*工序号"))
        key = exec_op_key_by_route_opno.get((route_code, op_no))
        if not key:
            continue
        grouped.setdefault(key, []).append(row)
    return grouped


def build_workcenter_map(factory_wb) -> dict[str, str]:
    result: dict[str, str] = {}
    for row in read_table(factory_wb["工作中心"]):
        name = normalize(row.get("*名称"))
        code = normalize(row.get("*编码"))
        if name and code and name not in result:
            result[name] = code
    return result


def build_source_routes(source_wb) -> dict[str, list[dict[str, str]]]:
    return {sheet_name: read_table(source_wb[sheet_name]) for sheet_name in SOURCE_ROUTE_SHEETS}


def build_new_route_data(
    source_routes: dict[str, list[dict[str, str]]],
    route_rows_by_name: dict[str, dict[str, str]],
    existing_exec_ops: dict[tuple[str, str], dict[str, str]],
    existing_op_materials: dict[tuple[str, str], list[dict[str, str]]],
    workcenter_codes: dict[str, str],
    material_codes: dict[str, str],
) -> dict[str, list[list[str]]]:
    route_op_rows: list[list[str]] = []
    route_seq_rows: list[list[str]] = []
    route_material_rows: list[list[str]] = []
    route_step_rows: list[list[str]] = []

    for route_name, source_rows in source_routes.items():
        route_meta = route_rows_by_name[route_name]
        route_code = normalize(route_meta.get("*编码"))
        route_version = normalize(route_meta.get("*版本号"))
        default_material_code = normalize(route_meta.get("物料编码"))
        previous_no = ""

        for source_row in source_rows:
            source_op_no = normalize(source_row.get("工序序号"))
            source_op_name = normalize(source_row.get("工序名称"))
            existing_exec = existing_exec_ops.get((route_name, source_op_name), {})
            old_exec_op_no = normalize(existing_exec.get("*工序号"))
            workcenter_code = normalize(existing_exec.get("*工作中心编码")) or workcenter_codes.get(
                normalize(source_row.get("工作中心")), ""
            )
            output_material_code = normalize(existing_exec.get("产出物料编码")) or default_material_code
            master_org = normalize(existing_exec.get("主制单位")) or normalize(route_meta.get("*工厂组织"))
            op_spec = normalize(existing_exec.get("工序专业类型")) or "装配专业"
            record_text = normalize(source_row.get("质量记录/防错")) or normalize(source_row.get("检测/记录要求"))
            key_control = normalize(source_row.get("关键控制点"))

            route_op_rows.append(
                [
                    source_op_no,
                    normalize(existing_exec.get("*工序类型")) or op_type(source_op_name),
                    route_content(source_row),
                    workcenter_code,
                    master_org,
                    normalize(existing_exec.get("产出物料版本号")) or route_version,
                    output_material_code,
                    "0",
                    normalize(source_row.get("工时（分钟）")),
                    normalize(existing_exec.get("*时间单位")) or "分钟",
                    normalize(existing_exec.get("执行标记")) or "是",
                    normalize(existing_exec.get("产出比")) or "1",
                    route_version,
                    route_code,
                    source_op_name,
                    previous_no,
                    op_spec,
                ]
            )

            if previous_no:
                route_seq_rows.append(["ES", source_op_no, previous_no, route_version, route_code])

            for material_row in existing_op_materials.get((route_name, source_op_name), []):
                route_material_rows.append(
                    [
                        route_version,
                        route_code,
                        normalize(material_row.get("*物料版本号")) or route_version,
                        normalize(material_row.get("*物料编码")),
                        source_op_no,
                        normalize(material_row.get("数量")),
                    ]
                )

            if key_control:
                route_step_rows.append([route_version, route_code, source_op_no, "10", "关键控制", key_control])
            if record_text:
                route_step_rows.append([route_version, route_code, source_op_no, "20", "质量记录", record_text])

            previous_no = source_op_no

    return {
        "工艺路线工序": route_op_rows,
        "工艺路线工序序列": route_seq_rows,
        "工艺路线工序物料": route_material_rows,
        "工艺路线工步": route_step_rows,
    }


def parse_args():
    parser = argparse.ArgumentParser(description="按样例工艺数据更新产品与工艺模板，并沿用已有编码")
    parser.add_argument("--source", type=Path, required=True, help="样例来源 Excel")
    parser.add_argument("--product", type=Path, required=True, help="原产品与工艺模板")
    parser.add_argument("--factory", type=Path, required=True, help="工厂资源模板，用于工作中心编码映射")
    parser.add_argument("--output", type=Path, required=True, help="输出的新产品与工艺模板")
    return parser.parse_args()


def main() -> int:
    args = parse_args()

    shutil.copy2(args.product, args.output)

    product_wb = load_workbook(args.product, data_only=False)
    factory_wb = load_workbook(args.factory, data_only=True)
    source_wb = load_workbook(args.source, data_only=True)
    output_wb = load_workbook(args.output)

    try:
        material_codes = build_material_code_map(product_wb)
        route_rows_by_name = build_route_rows_by_name(product_wb)
        existing_exec_ops = build_existing_exec_op_map(product_wb)
        existing_op_materials = build_old_op_materials_map(product_wb)
        workcenter_codes = build_workcenter_map(factory_wb)
        source_routes = build_source_routes(source_wb)

        _ = material_codes  # reserved for future fallback mappings
        new_route_data = build_new_route_data(
            source_routes,
            route_rows_by_name,
            existing_exec_ops,
            existing_op_materials,
            workcenter_codes,
            material_codes,
        )

        original_output = load_workbook(args.output, data_only=False)
        original_non_route = {
            sheet_name: snapshot_rows(original_output[sheet_name])
            for sheet_name in original_output.sheetnames
            if sheet_name not in ROUTE_SHEETS_TO_REWRITE
        }

        for sheet_name in ROUTE_SHEETS_TO_REWRITE:
            clear_and_write_rows(output_wb[sheet_name], new_route_data[sheet_name])

        output_wb.save(args.output)

        verify_wb = load_workbook(args.output, data_only=False)
        for sheet_name, rows in original_non_route.items():
            if snapshot_rows(verify_wb[sheet_name]) != rows:
                raise RuntimeError(f"非工艺路线 sheet 被意外修改: {sheet_name}")
        for sheet_name in ROUTE_SHEETS_TO_REWRITE:
            actual = snapshot_rows(verify_wb[sheet_name])[1:]
            expected = [tuple(normalize(cell) for cell in row) for row in new_route_data[sheet_name]]
            if actual != expected:
                raise RuntimeError(f"工艺路线 sheet 写入校验失败: {sheet_name}")

        print(f"[OK] 已生成: {args.output.resolve()}")
        for sheet_name in ROUTE_SHEETS_TO_REWRITE:
            print(f"[OK] {sheet_name}: {len(new_route_data[sheet_name])} rows")
        return 0
    finally:
        product_wb.close()
        factory_wb.close()
        source_wb.close()
        output_wb.close()


if __name__ == "__main__":
    raise SystemExit(main())
