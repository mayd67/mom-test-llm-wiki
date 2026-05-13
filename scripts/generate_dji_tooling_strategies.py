from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass


DEFAULT_FACTORY_ORG = "BizDJ008"
DEFAULT_EXEC_ORG = "BizDJ006"

SH_TOOLING = "工装工具"
SH_INSP_STRATEGY = "工装检定策略"
SH_INSP_ITEM = "工装检定项"
SH_MAINT_STRATEGY = "工装保养策略"
SH_MAINT_SPARE = "工装保养备件明细"
SH_MAINT_ITEM = "工装保养项"
SH_SOURCE_TOOL = "工具工装台帐表"


def normalize(value) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    return str(value).strip()


def read_table(ws) -> list[dict[str, str]]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [normalize(cell) for cell in rows[0]]
    result: list[dict[str, str]] = []
    for raw_row in rows[1:]:
        item: dict[str, str] = {}
        has_value = False
        for idx, header in enumerate(headers):
            if not header:
                continue
            value = normalize(raw_row[idx] if idx < len(raw_row) else "")
            if value:
                has_value = True
            item[header] = value
        if has_value:
            result.append(item)
    return result


def clear_and_write_rows(ws, rows: list[list[str]]) -> None:
    max_col = max(ws.max_column, max((len(row) for row in rows), default=0))
    for r in range(2, ws.max_row + 1):
        for c in range(1, max_col + 1):
            ws.cell(r, c).value = None
    for r_idx, row in enumerate(rows, start=2):
        for c_idx, value in enumerate(row, start=1):
            ws.cell(r_idx, c_idx).value = value


def snapshot_rows(ws) -> list[tuple[str, ...]]:
    data: list[tuple[str, ...]] = []
    for row in ws.iter_rows(values_only=True):
        values = tuple(normalize(cell) for cell in row)
        if any(values):
            data.append(values)
    return data


def parse_cycle(text: str) -> tuple[str, str] | None:
    value = normalize(text)
    if not value or value in {"不适用", "/"}:
        return None
    month_match = re.search(r"(\d+)\s*个?月", value)
    if month_match:
        return "月", month_match.group(1)
    week_match = re.search(r"(\d+)\s*周", value)
    if week_match:
        return "周", week_match.group(1)
    day_match = re.search(r"(\d+)\s*天", value)
    if day_match:
        return "天", day_match.group(1)
    hour_match = re.search(r"(\d+)\s*小时", value)
    if hour_match:
        return "小时", hour_match.group(1)
    return None


def choose_advance(unit: str, value: int) -> str:
    if unit == "月":
        if value >= 24:
            return "30"
        if value >= 12:
            return "15"
        if value >= 6:
            return "7"
        if value >= 3:
            return "5"
        return "3"
    if unit == "周":
        return str(max(1, min(7, value)))
    if unit == "天":
        return str(max(1, min(7, value // 5 if value > 5 else 1)))
    if unit == "小时":
        return "1"
    return "3"


def choose_inspection_duration(tool_name: str, usage: str) -> str:
    text = f"{tool_name}{usage}"
    if any(keyword in text for keyword in ("回流焊炉", "AOI", "ICT", "云台稳定调试台", "频谱分析仪")):
        return "4"
    if any(keyword in text for keyword in ("示波器", "光功率计", "扭矩", "表面电阻", "温湿度计", "动平衡")):
        return "2"
    return "1"


def choose_maintenance_duration(tool_name: str, category: str) -> str:
    text = f"{tool_name}{category}"
    if any(keyword in text for keyword in ("回流焊炉", "AOI", "ICT", "云台稳定调试台", "专用工装")):
        return "2"
    return "1"


def choose_maintenance_cycle(life_days_text: str, management: str, inspection_cycle: tuple[str, str] | None) -> tuple[str, str]:
    life_days = normalize(life_days_text)
    if life_days.isdigit():
        days = int(life_days)
        if days % 30 == 0 and days >= 30:
            return "月", str(days // 30)
        return "天", str(days)

    mgmt = normalize(management)
    if "每季度" in mgmt:
        return "月", "3"
    if any(token in mgmt for token in ("每月", "日常", "每天", "每次")):
        return "月", "1"
    if "每年" in mgmt:
        return "月", "12"
    if inspection_cycle is not None:
        unit, value = inspection_cycle
        if unit == "月":
            if int(value) >= 12:
                return "月", "6"
            if int(value) >= 6:
                return "月", "3"
        if unit == "周":
            return "周", "2"
        if unit == "天":
            return "天", max("7", value)
    return "月", "3"


def inspection_code(tool_code: str) -> str:
    return tool_code.replace("TOL-", "JD-", 1) if tool_code.startswith("TOL-") else f"JD-{tool_code}"


def maintenance_code(tool_code: str) -> str:
    return tool_code.replace("TOL-", "BY-", 1) if tool_code.startswith("TOL-") else f"BY-{tool_code}"


def build_inspection_items(strategy_code: str, precision: str, usage: str, management: str, recent_date: str) -> list[list[str]]:
    items = [
        [strategy_code, "外观与标识检查", f"工装外观完好、标识清晰、校准状态可识别；管理要求：{management or '按标准执行'}"],
        [strategy_code, "精度/功能检定", f"满足精度要求：{precision or '按铭牌要求'}；能够支持{usage or '既定用途'}。"],
    ]
    if management or recent_date:
        items.append(
            [
                strategy_code,
                "记录与状态复核",
                f"核对最近检定记录：{recent_date or '无'}；执行要求：{management or '记录完整、状态可追溯'}。",
            ]
        )
    return items


def build_maintenance_items(strategy_code: str, tool_name: str, usage: str, management: str) -> list[list[str]]:
    return [
        [strategy_code, "清洁与外观检查", "清除表面灰尘、焊渣、油污和残留物，确认壳体、刻度、铭牌、保护件完好。", "15"],
        [strategy_code, "连接与紧固检查", f"检查电源线、探头、夹头、连接器及紧固部位，确保可满足{usage or tool_name}使用要求。", "20"],
        [strategy_code, "功能恢复与防护确认", f"执行保养后功能自检，落实专项要求：{management or '按工装日常管理要求执行'}。", "15"],
    ]


def strategy_remark(tool_code: str, drawing: str, source_cycle: str, management: str, precision: str) -> str:
    parts = [f"来源工装:{tool_code}"]
    if drawing:
        parts.append(f"图号:{drawing}")
    if precision:
        parts.append(f"精度要求:{precision}")
    if source_cycle:
        parts.append(f"源周期:{source_cycle}")
    if management:
        parts.append(f"管理要求:{management}")
    return "；".join(parts)


def write_rows_to_template(workbook_path: Path, sheet_rows: dict[str, list[list[str]]]) -> Path:
    wb = load_workbook(workbook_path)
    try:
        untouched = {
            ws.title: snapshot_rows(ws)
            for ws in wb.worksheets
            if ws.title not in sheet_rows
        }
        for sheet_name, rows in sheet_rows.items():
            clear_and_write_rows(wb[sheet_name], rows)

        actual_output = workbook_path
        try:
            wb.save(actual_output)
        except PermissionError:
            actual_output = workbook_path.with_name(f"{workbook_path.stem}_已生成{workbook_path.suffix}")
            wb.save(actual_output)

        verify = load_workbook(actual_output, data_only=True)
        try:
            for name, rows in untouched.items():
                if snapshot_rows(verify[name]) != rows:
                    raise RuntimeError(f"非目标sheet发生变化: {name}")
            for sheet_name, rows in sheet_rows.items():
                actual = snapshot_rows(verify[sheet_name])[1:]
                expected = [tuple(normalize(cell) for cell in row) for row in rows]
                if actual != expected:
                    raise RuntimeError(f"sheet写入校验失败: {sheet_name}")
        finally:
            verify.close()

        return actual_output
    finally:
        wb.close()


def parse_args():
    parser = argparse.ArgumentParser(description="根据大疆无人机工装工具数据生成检定/保养策略模板")
    parser.add_argument("--resource-workbook", type=Path, required=True, help="工厂资源模板")
    parser.add_argument("--source-workbook", type=Path, required=True, help="原始大疆无人机BOM与工艺路线工作簿")
    parser.add_argument("--inspection-template", type=Path, required=True, help="工装检定策略模板")
    parser.add_argument("--maintenance-template", type=Path, required=True, help="工装保养策略模板")
    return parser.parse_args()


def main() -> int:
    args = parse_args()

    resource_wb = load_workbook(args.resource_workbook.resolve(), data_only=True)
    source_wb = load_workbook(args.source_workbook.resolve(), data_only=True)
    try:
        tooling_rows = read_table(resource_wb[SH_TOOLING])
        source_tool_rows = read_table(source_wb[SH_SOURCE_TOOL])
    finally:
        resource_wb.close()
        source_wb.close()

    source_by_drawing = {row.get("工具编码", ""): row for row in source_tool_rows if row.get("工具编码")}

    inspection_strategy_rows: list[list[str]] = []
    inspection_item_rows: list[list[str]] = []
    maintenance_strategy_rows: list[list[str]] = []
    maintenance_spare_rows: list[list[str]] = []
    maintenance_item_rows: list[list[str]] = []

    for tool in tooling_rows:
        tool_name = normalize(tool.get("*名称"))
        drawing = normalize(tool.get("图号"))
        tool_code = normalize(tool.get("*编码"))
        category = normalize(tool.get("工装类别")) or "通用工具"
        version = normalize(tool.get("*版本号")) or "A.01"
        security = normalize(tool.get("*密级")) or "公开"
        life_days = normalize(tool.get("理论寿命(天)"))
        source_tool = source_by_drawing.get(drawing, {})
        usage = normalize(source_tool.get("用途"))
        precision = normalize(source_tool.get("精度要求"))
        cycle_text = normalize(source_tool.get("定检周期"))
        recent_date = normalize(source_tool.get("最近检定日期（示例）"))
        management = normalize(source_tool.get("管理要求"))

        inspection_cycle = parse_cycle(cycle_text)
        if inspection_cycle is not None:
            cycle_unit, cycle_value = inspection_cycle
            insp_code = inspection_code(tool_code)
            inspection_strategy_rows.append(
                [
                    f"{tool_name}检定策略",
                    category,
                    cycle_unit,
                    cycle_value,
                    choose_advance(cycle_unit, int(cycle_value)),
                    "是",
                    DEFAULT_EXEC_ORG,
                    choose_inspection_duration(tool_name, usage),
                    insp_code,
                    strategy_remark(tool_code, drawing, cycle_text, management, precision),
                    security,
                    DEFAULT_FACTORY_ORG,
                ]
            )
            inspection_item_rows.extend(
                build_inspection_items(insp_code, precision, usage, management, recent_date)
            )

        maint_unit, maint_value = choose_maintenance_cycle(life_days, management, inspection_cycle)
        maint_code = maintenance_code(tool_code)
        maintenance_strategy_rows.append(
            [
                f"{tool_name}保养策略",
                category,
                "时间周期",
                maint_unit,
                maint_value,
                choose_maintenance_duration(tool_name, category),
                "",
                choose_advance(maint_unit, int(maint_value)),
                "是",
                DEFAULT_EXEC_ORG,
                maint_code,
                strategy_remark(tool_code, drawing, f"理论寿命天数:{life_days}" if life_days else cycle_text, management, precision),
                security,
                DEFAULT_FACTORY_ORG,
            ]
        )
        maintenance_item_rows.extend(build_maintenance_items(maint_code, tool_name, usage, management))

    inspection_output = write_rows_to_template(
        args.inspection_template.resolve(),
        {
            SH_INSP_STRATEGY: inspection_strategy_rows,
            SH_INSP_ITEM: inspection_item_rows,
        },
    )
    maintenance_output = write_rows_to_template(
        args.maintenance_template.resolve(),
        {
            SH_MAINT_STRATEGY: maintenance_strategy_rows,
            SH_MAINT_SPARE: maintenance_spare_rows,
            SH_MAINT_ITEM: maintenance_item_rows,
        },
    )

    print(f"[OK] 检定策略文件: {inspection_output}")
    print(f"[OK] 检定策略: {len(inspection_strategy_rows)} 条, 检定项: {len(inspection_item_rows)} 条")
    print(f"[OK] 保养策略文件: {maintenance_output}")
    print(f"[OK] 保养策略: {len(maintenance_strategy_rows)} 条, 保养项: {len(maintenance_item_rows)} 条, 备件明细: {len(maintenance_spare_rows)} 条")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
