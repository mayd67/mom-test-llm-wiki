from __future__ import annotations

import argparse
import json
import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook

try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass

WB_SYSTEM = "\u7cfb\u7edf\u914d\u7f6e_\u6a21\u677f.xlsx"
WB_FACTORY = "\u5de5\u5382\u8d44\u6e90_\u6a21\u677f.xlsx"
WB_PRODUCT = "\u4ea7\u54c1\u4e0e\u5de5\u827a_\u6a21\u677f.xlsx"
WB_ORDER = "\u751f\u4ea7\u8ba2\u5355_\u6a21\u677f.xlsx"
TEMPLATE_FILES = (WB_SYSTEM, WB_FACTORY, WB_PRODUCT, WB_ORDER)
TEMPLATE_DIR_NAMES = ("templates", "模板")
REPO_MARKERS = ("raw", "wiki", "outputs", "skills")

SH_ADMIN = "\u884c\u653f\u7ec4\u7ec7"
SH_BIZ = "\u4e1a\u52a1\u7ec4\u7ec7"
SH_USER = "\u7528\u6237"
SH_SUPPLIER = "\u4f9b\u5e94\u5546"
SH_EQUIPMENT = "\u8bbe\u5907"
SH_EQUIP_USER = "\u8bbe\u5907\u4e0e\u7528\u6237\u7684\u5173\u7cfb\u5b9e\u4f53\u7c7b"
SH_TOOLING = "\u5de5\u88c5\u5de5\u5177"
SH_TOOLING_INSPECTION_REL = "\u5de5\u88c5\u68c0\u5b9a\u7b56\u7565\u5173\u7cfb"
SH_TOOLING_MAINT_REL = "\u5de5\u88c5\u4fdd\u517b\u7b56\u7565\u5173\u7cfb"
SH_WORKCENTER = "\u5de5\u4f5c\u4e2d\u5fc3"
SH_WC_USER = "\u5de5\u4f5c\u4e2d\u5fc3\u4e0e\u7528\u6237\u5173\u7cfb"
SH_WC_SUPPLIER = "\u5de5\u4f5c\u4e2d\u5fc3\u4e0e\u4f9b\u5e94\u5546\u7684\u5173\u7cfb"
SH_WC_EQUIPMENT = "\u5de5\u4f5c\u4e2d\u5fc3\u4e0e\u8bbe\u5907\u7684\u5173\u7cfb"
SH_WAREHOUSE = "\u5e93\u623f"
SH_LOCATION = "\u5e93\u4f4d"
SH_PROCESS_LIB = "\u5de5\u5e8f\u5e93"
SH_MATERIAL = "\u7269\u6599"
SH_MBOM = "MBOM"
SH_MBOM_NODE = "MBOM\u8282\u70b9"
SH_ROUTE = "\u5de5\u827a\u8def\u7ebf"
SH_ROUTE_OP = "\u5de5\u827a\u8def\u7ebf\u5de5\u5e8f"
SH_ROUTE_SEQ = "\u5de5\u827a\u8def\u7ebf\u5de5\u5e8f\u5e8f\u5217"
SH_ROUTE_MAT = "\u5de5\u827a\u8def\u7ebf\u5de5\u5e8f\u7269\u6599"
SH_ROUTE_STEP = "\u5de5\u827a\u8def\u7ebf\u5de5\u6b65"
SH_ORDER = "\u751f\u4ea7\u8ba2\u5355"
SH_PICK = "\u5907\u6599\u6e05\u5355"

COL_CODE = "*\u7f16\u7801"
COL_VERSION = "*\u7248\u672c\u53f7"
COL_PARENT_ORG = "*\u7236\u7ec4\u7ec7\u7f16\u7801"
COL_ADMIN_CODE = "\u884c\u653f\u7ec4\u7ec7\u7f16\u7801"
COL_BIZ_CODE = "\u4e1a\u52a1\u7ec4\u7ec7\u7f16\u7801"
COL_ORG_TYPE = "\u5de5\u5382\u7ec4\u7ec7\u7c7b\u578b"
COL_USER_ID = "*\u7f16\u53f7"
COL_FACTORY_ORG = "*\u5de5\u5382\u7ec4\u7ec7"
COL_TOOLING_VERSION = "*\u5de5\u88c5\u5de5\u5177\u7248\u672c\u53f7"
COL_TOOLING_CODE = "*\u5de5\u88c5\u5de5\u5177\u7f16\u7801"
COL_INSPECTION_POLICY_CODE = "*\u68c0\u5b9a\u7b56\u7565\u7f16\u7801"
COL_MAINT_POLICY_CODE = "*\u4fdd\u517b\u7b56\u7565\u7f16\u7801"
COL_DEVICE_CODE = "*\u8bbe\u5907\u7f16\u7801"
COL_USER_REF = "*\u7528\u6237"
COL_WC_CODE = "*\u5de5\u4f5c\u4e2d\u5fc3\u7f16\u7801"
COL_SUPPLIER = "*\u4f9b\u5e94\u5546"
COL_WAREHOUSE_CODE = "*\u5e93\u623f\u7f16\u7801"
COL_MATERIAL_CODE = "*\u7269\u6599\u7f16\u7801"
COL_MATERIAL_VERSION = "*\u7269\u6599\u7248\u672c\u53f7"
COL_MBOM_CODE = "*MBom\u7f16\u7801"
COL_MBOM_VERSION = "*MBom\u7248\u672c\u53f7"
COL_NODE_LEVEL = "*\u5c42\u7ea7"
COL_NODE_SEQ = "*\u5e8f\u53f7"
COL_PARENT_MATERIAL_CODE = "\u7236\u7269\u6599\u7f16\u7801"
COL_PARENT_MATERIAL_VERSION = "\u7236\u7269\u6599\u7248\u672c"
COL_NODE_MATERIAL_VERSION = "*\u7269\u6599\u7248\u672c"
COL_ROUTE_CODE = "*\u5de5\u827a\u8def\u7ebf\u7f16\u7801"
COL_ROUTE_VERSION = "*\u5de5\u827a\u8def\u7ebf\u7248\u672c\u53f7"
COL_OP_NO = "*\u5de5\u5e8f\u53f7"
COL_PREV_OP = "\u524d\u7f6e\u5de5\u5e8f"
COL_OUTPUT_MATERIAL_CODE = "\u4ea7\u51fa\u7269\u6599\u7f16\u7801"
COL_OUTPUT_MATERIAL_VERSION = "\u4ea7\u51fa\u7269\u6599\u7248\u672c\u53f7"
COL_UPSTREAM_OP = "*\u4e0a\u9053\u5de5\u5e8f\u53f7"
COL_STEP_NO = "*\u5de5\u6b65\u5e8f\u53f7"
COL_ORDER_CODE = "*\u751f\u4ea7\u8ba2\u5355\u7f16\u7801"
COL_PLANNER = "\u8ba1\u5212\u5458"
COL_ORDER_ORG = "*\u6240\u5c5e\u7ec4\u7ec7"
COL_ORDER_BOM_VERSION = "BOM\u7248\u672c\u53f7"
COL_ORDER_BOM_CODE = "BOM\u7f16\u7801"
COL_ORDER_ROUTE_VERSION = "\u5de5\u827a\u8def\u7ebf\u7248\u672c\u53f7"
COL_ORDER_ROUTE_CODE = "\u5de5\u827a\u8def\u7ebf\u7f16\u7801"
COL_ROUTE_OP_CODE = "\u5de5\u5e8f\u7f16\u7801"
COL_ROUTE_OP_NAME = "\u5de5\u5e8f\u540d\u79f0"
COL_REPL_MATERIAL_VERSION = "\u66ff\u6362\u4ef6\u7269\u6599\u7248\u672c\u53f7"
COL_REPL_MATERIAL_CODE = "\u66ff\u6362\u4ef6\u7269\u6599\u7f16\u7801"

UNIQUE_KEYS = {
    (WB_SYSTEM, SH_ADMIN): (COL_CODE,),
    (WB_SYSTEM, SH_BIZ): (COL_CODE,),
    (WB_SYSTEM, SH_USER): (COL_USER_ID,),
    (WB_FACTORY, SH_SUPPLIER): (COL_CODE,),
    (WB_FACTORY, SH_EQUIPMENT): (COL_CODE,),
    (WB_FACTORY, SH_EQUIP_USER): (COL_DEVICE_CODE, COL_USER_REF),
    (WB_FACTORY, SH_TOOLING): (COL_CODE,),
    (WB_FACTORY, SH_TOOLING_INSPECTION_REL): (COL_TOOLING_VERSION, COL_TOOLING_CODE, COL_INSPECTION_POLICY_CODE, COL_FACTORY_ORG),
    (WB_FACTORY, SH_TOOLING_MAINT_REL): (COL_TOOLING_VERSION, COL_TOOLING_CODE, COL_MAINT_POLICY_CODE, COL_FACTORY_ORG),
    (WB_FACTORY, SH_WORKCENTER): (COL_CODE,),
    (WB_FACTORY, SH_WC_USER): (COL_WC_CODE, COL_USER_REF),
    (WB_FACTORY, SH_WC_SUPPLIER): (COL_WC_CODE, COL_SUPPLIER),
    (WB_FACTORY, SH_WC_EQUIPMENT): (COL_DEVICE_CODE, COL_WC_CODE),
    (WB_FACTORY, SH_WAREHOUSE): (COL_CODE,),
    (WB_FACTORY, SH_LOCATION): (COL_CODE,),
    (WB_FACTORY, SH_PROCESS_LIB): ("*\u540d\u79f0", COL_WC_CODE),
    (WB_PRODUCT, SH_MATERIAL): (COL_CODE,),
    (WB_PRODUCT, SH_MBOM): (COL_CODE,),
    (WB_PRODUCT, SH_MBOM_NODE): (COL_MBOM_CODE, COL_NODE_LEVEL, COL_NODE_SEQ),
    (WB_PRODUCT, SH_ROUTE): (COL_CODE,),
    (WB_PRODUCT, SH_ROUTE_OP): (COL_ROUTE_CODE, COL_OP_NO),
    (WB_PRODUCT, SH_ROUTE_SEQ): (COL_ROUTE_CODE, COL_OP_NO, COL_UPSTREAM_OP),
    (WB_PRODUCT, SH_ROUTE_MAT): (COL_ROUTE_CODE, COL_OP_NO, COL_MATERIAL_CODE),
    (WB_PRODUCT, SH_ROUTE_STEP): (COL_ROUTE_CODE, COL_OP_NO, COL_STEP_NO),
    (WB_ORDER, SH_ORDER): (COL_CODE,),
    (WB_ORDER, SH_PICK): (COL_ORDER_CODE, COL_MATERIAL_VERSION, COL_MATERIAL_CODE, COL_ROUTE_OP_CODE),
}

REFERENCE_RULES = [
    ((WB_SYSTEM, SH_ADMIN, (COL_PARENT_ORG,)), (WB_SYSTEM, SH_ADMIN, (COL_CODE,)), True),
    ((WB_SYSTEM, SH_BIZ, (COL_PARENT_ORG,)), (WB_SYSTEM, SH_BIZ, (COL_CODE,)), True),
    ((WB_SYSTEM, SH_BIZ, (COL_ADMIN_CODE,)), (WB_SYSTEM, SH_ADMIN, (COL_CODE,)), False),
    ((WB_SYSTEM, SH_USER, (COL_ADMIN_CODE,)), (WB_SYSTEM, SH_ADMIN, (COL_CODE,)), False),
    ((WB_SYSTEM, SH_USER, (COL_BIZ_CODE,)), (WB_SYSTEM, SH_BIZ, (COL_CODE,)), False),
    ((WB_FACTORY, SH_EQUIPMENT, (COL_FACTORY_ORG,)), (WB_SYSTEM, SH_BIZ, (COL_CODE,)), False),
    ((WB_FACTORY, SH_EQUIP_USER, (COL_DEVICE_CODE,)), (WB_FACTORY, SH_EQUIPMENT, (COL_CODE,)), False),
    ((WB_FACTORY, SH_EQUIP_USER, (COL_USER_REF,)), (WB_SYSTEM, SH_USER, (COL_USER_ID,)), False),
    ((WB_FACTORY, SH_TOOLING_INSPECTION_REL, (COL_TOOLING_VERSION, COL_TOOLING_CODE)), (WB_FACTORY, SH_TOOLING, (COL_VERSION, COL_CODE)), False),
    ((WB_FACTORY, SH_TOOLING_INSPECTION_REL, (COL_FACTORY_ORG,)), (WB_SYSTEM, SH_BIZ, (COL_CODE,)), False),
    ((WB_FACTORY, SH_TOOLING_MAINT_REL, (COL_TOOLING_VERSION, COL_TOOLING_CODE)), (WB_FACTORY, SH_TOOLING, (COL_VERSION, COL_CODE)), False),
    ((WB_FACTORY, SH_TOOLING_MAINT_REL, (COL_FACTORY_ORG,)), (WB_SYSTEM, SH_BIZ, (COL_CODE,)), False),
    ((WB_FACTORY, SH_WORKCENTER, (COL_FACTORY_ORG,)), (WB_SYSTEM, SH_BIZ, (COL_CODE,)), False),
    ((WB_FACTORY, SH_WC_USER, (COL_WC_CODE,)), (WB_FACTORY, SH_WORKCENTER, (COL_CODE,)), False),
    ((WB_FACTORY, SH_WC_USER, (COL_USER_REF,)), (WB_SYSTEM, SH_USER, (COL_USER_ID,)), False),
    ((WB_FACTORY, SH_WC_SUPPLIER, (COL_WC_CODE,)), (WB_FACTORY, SH_WORKCENTER, (COL_CODE,)), False),
    ((WB_FACTORY, SH_WC_SUPPLIER, (COL_SUPPLIER,)), (WB_FACTORY, SH_SUPPLIER, (COL_CODE,)), False),
    ((WB_FACTORY, SH_WC_EQUIPMENT, (COL_DEVICE_CODE,)), (WB_FACTORY, SH_EQUIPMENT, (COL_CODE,)), False),
    ((WB_FACTORY, SH_WC_EQUIPMENT, (COL_WC_CODE,)), (WB_FACTORY, SH_WORKCENTER, (COL_CODE,)), False),
    ((WB_FACTORY, SH_WAREHOUSE, (COL_FACTORY_ORG,)), (WB_SYSTEM, SH_BIZ, (COL_CODE,)), False),
    ((WB_FACTORY, SH_LOCATION, (COL_FACTORY_ORG,)), (WB_SYSTEM, SH_BIZ, (COL_CODE,)), False),
    ((WB_FACTORY, SH_LOCATION, (COL_WAREHOUSE_CODE,)), (WB_FACTORY, SH_WAREHOUSE, (COL_CODE,)), False),
    ((WB_FACTORY, SH_PROCESS_LIB, (COL_WC_CODE,)), (WB_FACTORY, SH_WORKCENTER, (COL_CODE,)), False),
    ((WB_PRODUCT, SH_MBOM, (COL_MATERIAL_VERSION, COL_MATERIAL_CODE)), (WB_PRODUCT, SH_MATERIAL, (COL_VERSION, COL_CODE)), False),
    ((WB_PRODUCT, SH_MBOM_NODE, (COL_MBOM_VERSION, COL_MBOM_CODE)), (WB_PRODUCT, SH_MBOM, (COL_VERSION, COL_CODE)), False),
    ((WB_PRODUCT, SH_MBOM_NODE, (COL_NODE_MATERIAL_VERSION, COL_MATERIAL_CODE)), (WB_PRODUCT, SH_MATERIAL, (COL_VERSION, COL_CODE)), False),
    ((WB_PRODUCT, SH_MBOM_NODE, (COL_PARENT_MATERIAL_VERSION, COL_PARENT_MATERIAL_CODE)), (WB_PRODUCT, SH_MATERIAL, (COL_VERSION, COL_CODE)), False),
    ((WB_PRODUCT, SH_ROUTE, ("\u7269\u6599\u7248\u672c\u53f7", "\u7269\u6599\u7f16\u7801")), (WB_PRODUCT, SH_MATERIAL, (COL_VERSION, COL_CODE)), False),
    ((WB_PRODUCT, SH_ROUTE, (COL_FACTORY_ORG,)), (WB_SYSTEM, SH_BIZ, (COL_CODE,)), False),
    ((WB_PRODUCT, SH_ROUTE_OP, (COL_WC_CODE,)), (WB_FACTORY, SH_WORKCENTER, (COL_CODE,)), False),
    ((WB_PRODUCT, SH_ROUTE_OP, (COL_OUTPUT_MATERIAL_VERSION, COL_OUTPUT_MATERIAL_CODE)), (WB_PRODUCT, SH_MATERIAL, (COL_VERSION, COL_CODE)), False),
    ((WB_PRODUCT, SH_ROUTE_OP, (COL_ROUTE_VERSION, COL_ROUTE_CODE)), (WB_PRODUCT, SH_ROUTE, (COL_VERSION, COL_CODE)), False),
    ((WB_PRODUCT, SH_ROUTE_SEQ, (COL_ROUTE_VERSION, COL_ROUTE_CODE)), (WB_PRODUCT, SH_ROUTE, (COL_VERSION, COL_CODE)), False),
    ((WB_PRODUCT, SH_ROUTE_MAT, (COL_ROUTE_VERSION, COL_ROUTE_CODE)), (WB_PRODUCT, SH_ROUTE, (COL_VERSION, COL_CODE)), False),
    ((WB_PRODUCT, SH_ROUTE_MAT, (COL_MATERIAL_VERSION, COL_MATERIAL_CODE)), (WB_PRODUCT, SH_MATERIAL, (COL_VERSION, COL_CODE)), False),
    ((WB_PRODUCT, SH_ROUTE_STEP, (COL_ROUTE_VERSION, COL_ROUTE_CODE)), (WB_PRODUCT, SH_ROUTE, (COL_VERSION, COL_CODE)), False),
    ((WB_ORDER, SH_ORDER, (COL_MATERIAL_VERSION, COL_MATERIAL_CODE)), (WB_PRODUCT, SH_MATERIAL, (COL_VERSION, COL_CODE)), False),
    ((WB_ORDER, SH_ORDER, (COL_ORDER_BOM_VERSION, COL_ORDER_BOM_CODE)), (WB_PRODUCT, SH_MBOM, (COL_VERSION, COL_CODE)), False),
    ((WB_ORDER, SH_ORDER, (COL_ORDER_ROUTE_VERSION, COL_ORDER_ROUTE_CODE)), (WB_PRODUCT, SH_ROUTE, (COL_VERSION, COL_CODE)), False),
    ((WB_ORDER, SH_ORDER, (COL_PLANNER,)), (WB_SYSTEM, SH_USER, (COL_USER_ID,)), False),
    ((WB_ORDER, SH_ORDER, (COL_ORDER_ORG,)), (WB_SYSTEM, SH_BIZ, (COL_CODE,)), False),
    ((WB_ORDER, SH_PICK, (COL_ORDER_CODE,)), (WB_ORDER, SH_ORDER, (COL_CODE,)), False),
    ((WB_ORDER, SH_PICK, (COL_MATERIAL_VERSION, COL_MATERIAL_CODE)), (WB_PRODUCT, SH_MATERIAL, (COL_VERSION, COL_CODE)), False),
    ((WB_ORDER, SH_PICK, (COL_REPL_MATERIAL_VERSION, COL_REPL_MATERIAL_CODE)), (WB_PRODUCT, SH_MATERIAL, (COL_VERSION, COL_CODE)), False),
]

DEFAULT_INTEGRATION_SYSTEM = "MOM-SEED"
DEFAULT_SECURITY_VALUE = "\u516c\u5f00"
HEADER_ALIASES = {
    (WB_SYSTEM, SH_USER): {
        "*\u540d\u79f0": ("*\u540d\u79f0", "\u540d\u79f0"),
        "\u540d\u79f0": ("\u540d\u79f0", "*\u540d\u79f0"),
        "*\u7528\u6237\u5b89\u5168\u7b49\u7ea7": ("*\u7528\u6237\u5b89\u5168\u7b49\u7ea7", "\u7528\u6237\u5b89\u5168\u7b49\u7ea7"),
        "\u7528\u6237\u5b89\u5168\u7b49\u7ea7": ("\u7528\u6237\u5b89\u5168\u7b49\u7ea7", "*\u7528\u6237\u5b89\u5168\u7b49\u7ea7"),
    },
    (WB_FACTORY, SH_WAREHOUSE): {
        "*\u4f5c\u4e1a\u6a21\u5f0f": ("*\u4f5c\u4e1a\u6a21\u5f0f", "\u4f5c\u4e1a\u6a21\u5f0f"),
        "\u4f5c\u4e1a\u6a21\u5f0f": ("\u4f5c\u4e1a\u6a21\u5f0f", "*\u4f5c\u4e1a\u6a21\u5f0f"),
    },
    (WB_FACTORY, SH_PROCESS_LIB): {
        "\u5de5\u5e8f\u4e13\u4e1a\u7c7b\u578b": ("\u5de5\u5e8f\u4e13\u4e1a\u7c7b\u578b", "\u5e8f\u4e13\u4e1a\u7c7b\u578b"),
        "\u5e8f\u4e13\u4e1a\u7c7b\u578b": ("\u5e8f\u4e13\u4e1a\u7c7b\u578b", "\u5de5\u5e8f\u4e13\u4e1a\u7c7b\u578b"),
    },
    (WB_PRODUCT, SH_MBOM_NODE): {
        "*MBom\u7248\u672c\u53f7": ("*MBom\u7248\u672c\u53f7", "*MBOM\u7248\u672c\u53f7"),
        "*MBOM\u7248\u672c\u53f7": ("*MBOM\u7248\u672c\u53f7", "*MBom\u7248\u672c\u53f7"),
        "*MBom\u7f16\u7801": ("*MBom\u7f16\u7801", "*MBOM\u7f16\u7801"),
        "*MBOM\u7f16\u7801": ("*MBOM\u7f16\u7801", "*MBom\u7f16\u7801"),
    },
}
GENERIC_HEADER_ALIASES = {
    "\u96c6\u6210\u521b\u5efa\u65f6\u95f4": ("\u96c6\u6210\u521b\u5efa\u65f6\u95f4", "*\u96c6\u6210\u521b\u5efa\u65f6\u95f4"),
    "*\u96c6\u6210\u521b\u5efa\u65f6\u95f4": ("*\u96c6\u6210\u521b\u5efa\u65f6\u95f4", "\u96c6\u6210\u521b\u5efa\u65f6\u95f4"),
}


def normalize(value):
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    return str(value).strip()


def _primary_identifier(row):
    for field in (
        COL_CODE,
        COL_USER_ID,
        COL_DEVICE_CODE,
        COL_WC_CODE,
        COL_WAREHOUSE_CODE,
        COL_MATERIAL_CODE,
        COL_ORDER_CODE,
        COL_MBOM_CODE,
        "*MBOM\u7f16\u7801",
        COL_ROUTE_CODE,
    ):
        value = normalize(row.get(field, ""))
        if value:
            return value
    return ""


def _candidate_fields(workbook_name, sheet_name, header):
    sheet_aliases = HEADER_ALIASES.get((workbook_name, sheet_name), {})
    if header in sheet_aliases:
        return sheet_aliases[header]
    if header in GENERIC_HEADER_ALIASES:
        return GENERIC_HEADER_ALIASES[header]
    return (header,)


def _default_field_value(seed, row, header):
    if header == "*\u5bc6\u7ea7":
        return normalize(seed.get("metadata", {}).get("default_security")) or DEFAULT_SECURITY_VALUE
    if header == "\u96c6\u6210\u7cfb\u7edf":
        return normalize(row.get(header, "")) or DEFAULT_INTEGRATION_SYSTEM
    if header == "\u96c6\u6210\u6570\u636e\u4e3b\u952e":
        return normalize(row.get(header, "")) or _primary_identifier(row)
    return ""


def align_seed_to_template(seed, template_headers):
    workbooks = seed.get("workbooks", {})
    for (workbook_name, sheet_name), headers in template_headers.items():
        workbook = workbooks.get(workbook_name)
        if workbook is None or sheet_name not in workbook:
            continue
        aligned_rows = []
        for row in workbook.get(sheet_name, []):
            aligned = {}
            for header in headers:
                value = ""
                for source_field in _candidate_fields(workbook_name, sheet_name, header):
                    if source_field in row and normalize(row.get(source_field, "")):
                        value = row[source_field]
                        break
                if not normalize(value):
                    value = _default_field_value(seed, row, header)
                if header.startswith("*") or normalize(value):
                    aligned[header] = value
            aligned_rows.append(aligned)
        workbook[sheet_name] = aligned_rows
    return seed


def is_complete_template_dir(candidate: Path) -> bool:
    return candidate.exists() and all((candidate / name).exists() for name in TEMPLATE_FILES)


def find_template_dir(start: Path):
    for candidate in [start, *start.parents]:
        for dirname in TEMPLATE_DIR_NAMES:
            nested = candidate / dirname
            if is_complete_template_dir(nested):
                return nested
        if is_complete_template_dir(candidate):
            return candidate
    return None


def resolve_template_dir(template_dir: Path | None, start: Path) -> Path:
    if template_dir is not None:
        resolved = template_dir.resolve()
        return find_template_dir(resolved) or resolved
    return find_template_dir(start) or start.parent.resolve()


def detect_repo_root(start: Path) -> Path | None:
    for candidate in [start.resolve(), *start.resolve().parents]:
        if all((candidate / marker).exists() for marker in REPO_MARKERS):
            return candidate
    return None


def default_output_root_for_template_dir(template_dir: Path) -> Path:
    resolved = template_dir.resolve()
    repo_root = detect_repo_root(resolved)
    if repo_root is not None:
        return repo_root / 'outputs' / '06_测试数据'
    if resolved.name in TEMPLATE_DIR_NAMES:
        owner_dir = resolved.parent
        if owner_dir.name == 'mom-template-seed-generator':
            return owner_dir.parent / 'output'
        return owner_dir / 'output'
    return resolved / 'output'


def build_parser():
    script_dir = Path(__file__).resolve().parent
    skill_dir = script_dir.parent
    template_dir = resolve_template_dir(None, skill_dir)
    output_root = default_output_root_for_template_dir(template_dir)
    parser = argparse.ArgumentParser(description="根据 JSON 种子生成 MOM Excel 数据包")
    parser.add_argument("--seed", type=Path, default=skill_dir / "assets" / "automotive_engine_seed.json")
    parser.add_argument("--template-dir", type=Path, default=template_dir)
    parser.add_argument("--output-dir", type=Path, default=output_root / "automotive_engine_seed")
    parser.add_argument("--validate-only", action="store_true")
    parser.add_argument("--dump-default-seed", type=Path, default=None)
    return parser


def rows(seed, workbook_name, sheet_name):
    return seed.get("workbooks", {}).get(workbook_name, {}).get(sheet_name, [])


def tuple_for(row, columns):
    return tuple(normalize(row.get(column, "")) for column in columns)


def complete(values):
    return all(normalize(value) for value in values)


def load_template_headers(template_dir: Path):
    headers = {}
    for workbook_name in TEMPLATE_FILES:
        workbook = load_workbook(template_dir / workbook_name)
        try:
            for worksheet in workbook.worksheets:
                headers[(workbook_name, worksheet.title)] = [normalize(cell.value) for cell in worksheet[1]]
        finally:
            workbook.close()
    return headers


def parse_validation_list(formula):
    formula = normalize(formula)
    if formula.startswith('"') and formula.endswith('"'):
        return [normalize(item) for item in formula[1:-1].split(',') if normalize(item)]
    return []


def load_template_list_validations(template_dir: Path):
    validations = {}
    for workbook_name in TEMPLATE_FILES:
        workbook = load_workbook(template_dir / workbook_name)
        try:
            for worksheet in workbook.worksheets:
                headers = [normalize(cell.value) for cell in worksheet[1]]
                if not worksheet.data_validations:
                    continue
                for dv in worksheet.data_validations.dataValidation:
                    if dv.type != 'list':
                        continue
                    allowed_values = parse_validation_list(dv.formula1)
                    if not allowed_values:
                        continue
                    for cell_range in dv.sqref.ranges:
                        if cell_range.min_col != cell_range.max_col:
                            continue
                        header = headers[cell_range.min_col - 1]
                        if not header:
                            continue
                        validations.setdefault((workbook_name, worksheet.title, header), set()).update(allowed_values)
        finally:
            workbook.close()
    return validations


def validate_list_values(seed, template_validations):
    errors = []
    for workbook_name, sheets in seed.get('workbooks', {}).items():
        for sheet_name, sheet_rows in sheets.items():
            for row_index, row in enumerate(sheet_rows, start=2):
                for field, value in row.items():
                    allowed = template_validations.get((workbook_name, sheet_name, field))
                    value = normalize(value)
                    if not allowed or not value:
                        continue
                    if value not in allowed:
                        errors.append(f"{label(workbook_name, sheet_name, row_index)} 枚举值不合法: {field}={value!r}, 可选={sorted(allowed)}")
    return errors


def label(workbook_name, sheet_name, row_index):
    return f"{workbook_name}/{sheet_name} 第{row_index}行"


def validate_structure(seed, template_headers):
    errors = []
    warnings = []
    for workbook_name, sheets in seed.get("workbooks", {}).items():
        if workbook_name not in TEMPLATE_FILES:
            errors.append(f"不支持的工作簿: {workbook_name!r}")
            continue
        for sheet_name, sheet_rows in sheets.items():
            key = (workbook_name, sheet_name)
            if key not in template_headers:
                errors.append(f"模板 {workbook_name!r} 中不存在 Sheet: {sheet_name!r}")
                continue
            headers = template_headers[key]
            required = [header for header in headers if header.startswith("*")]
            allowed = set(headers)
            if len(sheet_rows) > 100:
                warnings.append(f"{workbook_name}/{sheet_name} 当前共有 {len(sheet_rows)} 行，建议控制在 100 行以内")
            for row_index, row in enumerate(sheet_rows, start=2):
                unknown = [field for field in row.keys() if field not in allowed]
                if unknown:
                    errors.append(f"{label(workbook_name, sheet_name, row_index)} 存在未知字段: {unknown}")
                missing = [field for field in required if not normalize(row.get(field, ""))]
                if missing:
                    errors.append(f"{label(workbook_name, sheet_name, row_index)} 缺少必填字段: {missing}")
    return errors, warnings


def build_lookup(seed, workbook_name, sheet_name, columns):
    lookup = set()
    for row in rows(seed, workbook_name, sheet_name):
        key = tuple_for(row, columns)
        if complete(key):
            lookup.add(key)
    return lookup


def validate_unique(seed):
    errors = []
    for (workbook_name, sheet_name), columns in UNIQUE_KEYS.items():
        seen = {}
        for row_index, row in enumerate(rows(seed, workbook_name, sheet_name), start=2):
            key = tuple_for(row, columns)
            if not complete(key):
                continue
            if key in seen:
                errors.append(f"{label(workbook_name, sheet_name, row_index)} 与 {seen[key]} 在 {columns} 上重复")
            else:
                seen[key] = label(workbook_name, sheet_name, row_index)
    return errors


def validate_references(seed):
    errors = []
    external_refs = {normalize(item) for item in seed.get("external_references", [])}
    cache = {}

    def lookup(workbook_name, sheet_name, columns):
        cache_key = (workbook_name, sheet_name, tuple(columns))
        if cache_key not in cache:
            cache[cache_key] = build_lookup(seed, workbook_name, sheet_name, columns)
        return cache[cache_key]

    for source, target, allow_external in REFERENCE_RULES:
        src_wb, src_sh, src_cols = source
        tgt_wb, tgt_sh, tgt_cols = target
        target_values = lookup(tgt_wb, tgt_sh, tgt_cols)
        for row_index, row in enumerate(rows(seed, src_wb, src_sh), start=2):
            key = tuple_for(row, src_cols)
            if not any(key) or not complete(key):
                continue
            if key in target_values:
                continue
            if allow_external and len(key) == 1 and key[0] in external_refs:
                continue
            errors.append(f"{label(src_wb, src_sh, row_index)} 引用不存在: {src_cols} -> {key}")
    return errors


def validate_storage_org_rules(seed):
    errors = []
    plant_type = '工厂'
    error_prefix = '库房/库位的*工厂组织必须引用工厂级业务组织'
    biz_type_map = {
        normalize(row.get(COL_CODE, '')): normalize(row.get(COL_ORG_TYPE, ''))
        for row in rows(seed, WB_SYSTEM, SH_BIZ)
    }
    for sheet_name in (SH_WAREHOUSE, SH_LOCATION):
        for row_index, row in enumerate(rows(seed, WB_FACTORY, sheet_name), start=2):
            factory_org = normalize(row.get(COL_FACTORY_ORG, ''))
            if not factory_org:
                continue
            if biz_type_map.get(factory_org) != plant_type:
                errors.append(f"{label(WB_FACTORY, sheet_name, row_index)} {error_prefix}: {factory_org}")
    return errors


def validate_route_refs(seed):
    errors = []
    op_lookup = build_lookup(seed, WB_PRODUCT, SH_ROUTE_OP, (COL_ROUTE_VERSION, COL_ROUTE_CODE, COL_OP_NO))
    for row_index, row in enumerate(rows(seed, WB_PRODUCT, SH_ROUTE_OP), start=2):
        previous = normalize(row.get(COL_PREV_OP, ""))
        if previous:
            key = (normalize(row.get(COL_ROUTE_VERSION, "")), normalize(row.get(COL_ROUTE_CODE, "")), previous)
            if key not in op_lookup:
                errors.append(f"{label(WB_PRODUCT, SH_ROUTE_OP, row_index)} 前置工序不存在: {previous}")
    for row_index, row in enumerate(rows(seed, WB_PRODUCT, SH_ROUTE_SEQ), start=2):
        route_version = normalize(row.get(COL_ROUTE_VERSION, ""))
        route_code = normalize(row.get(COL_ROUTE_CODE, ""))
        for field in (COL_OP_NO, COL_UPSTREAM_OP):
            key = (route_version, route_code, normalize(row.get(field, "")))
            if key not in op_lookup:
                errors.append(f"{label(WB_PRODUCT, SH_ROUTE_SEQ, row_index)} 字段 {field!r} 引用了不存在的工序: {key}")
    for sheet_name in (SH_ROUTE_MAT, SH_ROUTE_STEP):
        for row_index, row in enumerate(rows(seed, WB_PRODUCT, sheet_name), start=2):
            key = (normalize(row.get(COL_ROUTE_VERSION, "")), normalize(row.get(COL_ROUTE_CODE, "")), normalize(row.get(COL_OP_NO, "")))
            if key not in op_lookup:
                errors.append(f"{label(WB_PRODUCT, sheet_name, row_index)} 引用了不存在的工序: {key}")
    return errors


def validate_order_refs(seed):
    errors = []
    order_rows = rows(seed, WB_ORDER, SH_ORDER)
    pick_rows = rows(seed, WB_ORDER, SH_PICK)
    route_rows = rows(seed, WB_PRODUCT, SH_ROUTE)
    mbom_rows = rows(seed, WB_PRODUCT, SH_MBOM)
    op_rows = rows(seed, WB_PRODUCT, SH_ROUTE_OP)

    route_map = {(normalize(row.get(COL_VERSION, "")), normalize(row.get(COL_CODE, ""))): row for row in route_rows}
    mbom_map = {(normalize(row.get(COL_VERSION, "")), normalize(row.get(COL_CODE, ""))): row for row in mbom_rows}
    order_map = {normalize(row.get(COL_CODE, "")): row for row in order_rows}
    op_map = {}
    for row in op_rows:
        key = (normalize(row.get(COL_ROUTE_VERSION, "")), normalize(row.get(COL_ROUTE_CODE, "")), normalize(row.get(COL_OP_NO, "")))
        op_map[key] = normalize(row.get("*工序名称", ""))

    for row_index, row in enumerate(order_rows, start=2):
        material_key = (normalize(row.get(COL_MATERIAL_VERSION, "")), normalize(row.get(COL_MATERIAL_CODE, "")))
        route_key = (normalize(row.get(COL_ORDER_ROUTE_VERSION, "")), normalize(row.get(COL_ORDER_ROUTE_CODE, "")))
        if route_key in route_map:
            route_row = route_map[route_key]
            route_material_key = (normalize(route_row.get("物料版本号", "")), normalize(route_row.get("物料编码", "")))
            if route_material_key != material_key:
                errors.append(f"{label(WB_ORDER, SH_ORDER, row_index)} 工艺路线对应物料与订单物料不一致")
        mbom_key = (normalize(row.get(COL_ORDER_BOM_VERSION, "")), normalize(row.get(COL_ORDER_BOM_CODE, "")))
        if all(mbom_key) and mbom_key in mbom_map:
            mbom_row = mbom_map[mbom_key]
            mbom_material_key = (normalize(mbom_row.get(COL_MATERIAL_VERSION, "")), normalize(mbom_row.get(COL_MATERIAL_CODE, "")))
            if mbom_material_key != material_key:
                errors.append(f"{label(WB_ORDER, SH_ORDER, row_index)} BOM 对应物料与订单物料不一致")

    for row_index, row in enumerate(pick_rows, start=2):
        order_code = normalize(row.get(COL_ORDER_CODE, ""))
        order_row = order_map.get(order_code)
        if not order_row:
            continue
        op_code = normalize(row.get(COL_ROUTE_OP_CODE, ""))
        if not op_code:
            continue
        op_key = (normalize(order_row.get(COL_ORDER_ROUTE_VERSION, "")), normalize(order_row.get(COL_ORDER_ROUTE_CODE, "")), op_code)
        op_name = op_map.get(op_key)
        if not op_name:
            errors.append(f"{label(WB_ORDER, SH_PICK, row_index)} 引用的工序编码不存在或不属于订单工艺路线: {op_code}")
            continue
        current_name = normalize(row.get(COL_ROUTE_OP_NAME, ""))
        if current_name and current_name != op_name:
            errors.append(f"{label(WB_ORDER, SH_PICK, row_index)} 工序名称与工艺路线工序名称不一致")
    return errors


def fill_workbook(template_path: Path, output_path: Path, sheet_rows_map):
    shutil.copy2(template_path, output_path)
    workbook = load_workbook(output_path)
    row_counts = {}
    try:
        for worksheet in workbook.worksheets:
            data_rows = sheet_rows_map.get(worksheet.title, [])
            if not data_rows:
                continue
            headers = [normalize(cell.value) for cell in worksheet[1]]
            header_map = {header: index for index, header in enumerate(headers, start=1)}
            for row_index, row in enumerate(data_rows, start=2):
                for field, value in row.items():
                    col_index = header_map.get(field)
                    if col_index is not None:
                        worksheet.cell(row=row_index, column=col_index).value = value
            row_counts[worksheet.title] = len(data_rows)
        workbook.save(output_path)
    finally:
        workbook.close()
    return row_counts


def build_summary(seed):
    return {
        "元数据": seed.get("metadata", {}),
        "外部引用": seed.get("external_references", []),
        "工作簿行数": {
            workbook_name: {sheet_name: len(sheet_rows) for sheet_name, sheet_rows in sheets.items()}
            for workbook_name, sheets in seed.get("workbooks", {}).items()
        },
    }

SECURITY = "\u5185\u90e8"
SECURITY = "\u5185\u90e8"
YES = "\u662f"
NO = "\u5426"
ADMIN_COMPANY = "\u516c\u53f8"
ADMIN_PLANT = "\u5de5\u5382"
ADMIN_DEPT = "\u90e8\u95e8"
BIZ_COMPANY = "\u516c\u53f8"
BIZ_PLANT = "\u5de5\u5382"
BIZ_DEPT = "\u90e8\u95e8"
BIZ_WORKSHOP = "\u8f66\u95f4"
FACTORY_MACH = "\u673a\u68b0\u52a0\u5de5\u4e13\u4e1a"
FACTORY_ASM = "\u88c5\u914d\u4e13\u4e1a"
USER_CORE = "\u6838\u5fc3"
USER_IMPORTANT = "\u91cd\u8981"
USER_NORMAL = "\u4e00\u822c"
GENDER_M = "\u7537"
GENDER_F = "\u5973"
WC_LINE = "\u4ea7\u7ebf"
WC_GROUP = "\u8bbe\u5907\u7ec4"
WC_OUT = "\u5916\u59d4"
WC_PROCESS = "\u52a0\u5de5"
WC_CHECK = "\u68c0\u9a8c"
WH_NORMAL = "\u666e\u901a\u5e93\u623f"
WH_ERP1 = "ERP\u4e00\u7ea7\u5e93"
WH_ERP2 = "ERP\u4e8c\u7ea7\u5e93"
WH_SHOP2 = "\u8f66\u95f4\u4e8c\u7ea7\u5e93"
MAT_CLASS = "\u7269\u6599"
TOOL_CLASS = "\u5de5\u88c5\u5de5\u5177"
CAT_RAW = "\u539f\u6750\u6599"
CAT_AUX = "\u8f85\u52a9\u6750\u6599"
CAT_PART = "\u96f6\u90e8\u4ef6"
CAT_KIT = "\u914d\u5957\u4ef6"
CAT_SPARE = "\u8bbe\u5907\u5907\u4ef6"
MAKE_SELF = "\u81ea\u5236\u4ef6"
MAKE_BUY = "\u5916\u8d2d\u4ef6"
FEATURE_KEY = "\u5173\u952e\u4ef6"
FEATURE_IMPORTANT = "\u91cd\u8981\u4ef6"
FEATURE_NORMAL = "\u4e00\u822c\u4ef6"
TOOL_SPECIAL = "\u4e13\u7528\u5de5\u88c5"
TOOL_COMMON = "\u901a\u7528\u5de5\u5177"
ROUTE_FORMAL = "\u6b63\u5f0f\u5de5\u827a"
SPEC_MACH = "\u673a\u52a0"
SPEC_ASM = "\u88c5\u914d"
OP_PROCESS = "\u52a0\u5de5"
OP_CHECK = "\u68c0\u9a8c"
OP_OUT = "\u5916\u59d4"
TIME_MIN = "\u5206\u949f"
SEQ_SSEE = "ES"
STAGE_MP = "\u91cf\u4ea7"
UNIT_PC = "\u4e2a"


def material(code, name, category, make_type, feature, drawing):
    return {
        '*\u7269\u6599\u5206\u7c7b': MAT_CLASS,
        '*\u540d\u79f0': name,
        '\u7269\u6599\u7c7b\u522b': category,
        '\u56fe\u53f7': drawing,
        '*\u5236\u9020\u7c7b\u578b': make_type,
        '\u8ba1\u91cf\u5355\u4f4d': UNIT_PC,
        '\u7279\u6027\u5206\u7c7b': feature,
        '\u542f\u7528\u6279\u6b21\u6807\u8bb0': YES,
        '\u542f\u7528\u5e8f\u5217\u53f7\u6807\u8bb0': NO,
        '\u7269\u6599\u9636\u6bb5': STAGE_MP,
        '*\u7248\u672c\u53f7': 'A.01',
        '*\u7f16\u7801': code,
        '*\u5bc6\u7ea7': SECURITY,
    }


def tooling(code, name, category, make_type, tool_category):
    return {
        '*\u7269\u6599\u5206\u7c7b': TOOL_CLASS,
        '*\u540d\u79f0': name,
        '\u7269\u6599\u7c7b\u522b': category,
        '\u56fe\u53f7': code,
        '*\u5236\u9020\u7c7b\u578b': make_type,
        '\u8ba1\u91cf\u5355\u4f4d': UNIT_PC,
        '\u7279\u6027\u5206\u7c7b': FEATURE_NORMAL,
        '\u542f\u7528\u6279\u6b21\u6807\u8bb0': NO,
        '\u542f\u7528\u5e8f\u5217\u53f7\u6807\u8bb0': NO,
        '\u5de5\u88c5\u7c7b\u522b': tool_category,
        '\u4e00\u6b21\u6027\u5de5\u88c5\u6807\u8bb0': NO,
        '\u5355\u4ef6\u5de5\u88c5\u6807\u8bb0': NO,
        '\u7269\u6599\u9636\u6bb5': STAGE_MP,
        '*\u7248\u672c\u53f7': 'A.01',
        '*\u7f16\u7801': code,
        '*\u5bc6\u7ea7': SECURITY,
    }


def route(code, name, spec, factory_org, material_code):
    return {
        '*\u540d\u79f0': name,
        '*\u5de5\u827a\u7c7b\u578b': ROUTE_FORMAL,
        '\u5de5\u827a\u4e13\u4e1a': spec,
        '\u7269\u6599\u7248\u672c\u53f7': 'A.01',
        '\u7269\u6599\u7f16\u7801': material_code,
        '*\u7248\u672c\u53f7': 'A.01',
        '*\u7f16\u7801': code,
        '*\u5bc6\u7ea7': SECURITY,
        '*\u5de5\u5382\u7ec4\u7ec7': factory_org,
    }


def op(route_code, op_no, name, op_type, wc_code, prep, run, spec_type, output_code, previous=''):
    row = {
        '*\u5de5\u5e8f\u53f7': op_no,
        '*\u5de5\u5e8f\u7c7b\u578b': op_type,
        '\u5de5\u5e8f\u5185\u5bb9': name,
        '*\u5de5\u4f5c\u4e2d\u5fc3\u7f16\u7801': wc_code,
        '*\u5b9a\u989d\u8f85\u52a9\u5de5\u65f6': prep,
        '*\u5b9a\u989d\u52a0\u5de5\u65f6\u95f4': run,
        '*\u65f6\u95f4\u5355\u4f4d': TIME_MIN,
        '\u6267\u884c\u6807\u8bb0': YES,
        '\u4ea7\u51fa\u6bd4': 1,
        '*\u5de5\u827a\u8def\u7ebf\u7248\u672c\u53f7': 'A.01',
        '*\u5de5\u827a\u8def\u7ebf\u7f16\u7801': route_code,
        '*\u5de5\u5e8f\u540d\u79f0': name,
        '\u4ea7\u51fa\u7269\u6599\u7248\u672c\u53f7': 'A.01',
        '\u4ea7\u51fa\u7269\u6599\u7f16\u7801': output_code,
        '\u5de5\u5e8f\u4e13\u4e1a\u7c7b\u578b': spec_type,
    }
    if previous:
        row[COL_PREV_OP] = previous
    return row


def route_material(route_code, op_no, material_code, qty):
    return {
        '*\u5de5\u827a\u8def\u7ebf\u7248\u672c\u53f7': 'A.01',
        '*\u5de5\u827a\u8def\u7ebf\u7f16\u7801': route_code,
        '*\u7269\u6599\u7248\u672c\u53f7': 'A.01',
        '*\u7269\u6599\u7f16\u7801': material_code,
        '*\u5de5\u5e8f\u53f7': op_no,
        '\u6570\u91cf': qty,
    }


def route_step(route_code, op_no, step_no, name, content):
    return {
        '*\u5de5\u827a\u8def\u7ebf\u7248\u672c\u53f7': 'A.01',
        '*\u5de5\u827a\u8def\u7ebf\u7f16\u7801': route_code,
        '*\u5de5\u5e8f\u53f7': op_no,
        '*\u5de5\u6b65\u5e8f\u53f7': step_no,
        '*\u5de5\u6b65\u540d\u79f0': name,
        '\u5de5\u6b65\u5185\u5bb9': content,
    }


def sequence(route_code, op_no, previous):
    return {
        '*\u63a5\u7eed\u5173\u7cfb': SEQ_SSEE,
        '*\u5de5\u5e8f\u53f7': op_no,
        '*\u4e0a\u9053\u5de5\u5e8f\u53f7': previous,
        '*\u5de5\u827a\u8def\u7ebf\u7248\u672c\u53f7': 'A.01',
        '*\u5de5\u827a\u8def\u7ebf\u7f16\u7801': route_code,
    }


def process_lib(name, spec, op_type, wc_code, prep, run, content):
    return {
        '\u5e8f\u4e13\u4e1a\u7c7b\u578b': spec,
        '*\u540d\u79f0': name,
        '*\u5de5\u5e8f\u7c7b\u578b': op_type,
        '*\u5de5\u4f5c\u4e2d\u5fc3\u7f16\u7801': wc_code,
        '*\u5b9a\u989d\u51c6\u5907\u65f6\u95f4': prep,
        '*\u5b9a\u989d\u52a0\u5de5\u65f6\u95f4': run,
        '\u6267\u884c\u6807\u8bb0': YES,
        '*\u65f6\u95f4\u5355\u4f4d': TIME_MIN,
        '\u4ea7\u51fa\u6bd4': 1,
        '\u5de5\u5e8f\u5185\u5bb9': content,
        '*\u5bc6\u7ea7': SECURITY,
    }


def default_seed():
    asset_path = Path(__file__).resolve().parent.parent / "assets" / "automotive_engine_seed.json"
    if not asset_path.exists():
        raise FileNotFoundError(f"未找到默认种子文件: {asset_path}")
    return json.loads(asset_path.read_text(encoding="utf-8-sig"))

def write_default_seed(path: Path):
    seed = default_seed()
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(seed, ensure_ascii=False, indent=2), encoding='utf-8')


def main():
    parser = build_parser()
    args = parser.parse_args()
    if args.dump_default_seed:
        write_default_seed(args.dump_default_seed.resolve())
        print(f"[OK] 已导出默认种子: {args.dump_default_seed.resolve()}")
        return 0
    seed_path = args.seed.resolve()
    script_dir = Path(__file__).resolve().parent
    skill_dir = script_dir.parent
    template_dir = resolve_template_dir(args.template_dir, skill_dir)
    output_dir = args.output_dir.resolve()
    if not seed_path.exists():
        print(f"种子文件不存在: {seed_path}", file=sys.stderr)
        return 1
    if not is_complete_template_dir(template_dir):
        print(f"模板目录无效: {template_dir}", file=sys.stderr)
        return 1
    seed = json.loads(seed_path.read_text(encoding='utf-8-sig'))
    headers = load_template_headers(template_dir)
    align_seed_to_template(seed, headers)
    template_validations = load_template_list_validations(template_dir)
    structure_errors, warnings = validate_structure(seed, headers)
    errors = structure_errors + validate_list_values(seed, template_validations) + validate_unique(seed) + validate_references(seed) + validate_storage_org_rules(seed) + validate_route_refs(seed) + validate_order_refs(seed)
    for warning in warnings:
        print(f"[WARN] {warning}")
    if errors:
        print("[ERROR] 校验失败:")
        for error in errors:
            print(f"  - {error}")
        return 1
    summary = build_summary(seed)
    print("[OK] 校验通过")
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    if args.validate_only:
        return 0
    output_dir.mkdir(parents=True, exist_ok=True)
    for workbook_name in TEMPLATE_FILES:
        row_counts = fill_workbook(template_dir / workbook_name, output_dir / workbook_name, seed.get("workbooks", {}).get(workbook_name, {}))
        print(f"[OK] 已生成 {output_dir / workbook_name}")
        print(json.dumps(row_counts, ensure_ascii=False))
    legacy_summary = output_dir / "seed_summary.json"
    if legacy_summary.exists():
        legacy_summary.unlink()
    (output_dir / "种子概览.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
    print(f"[OK] 已写入摘要: {output_dir / '种子概览.json'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

