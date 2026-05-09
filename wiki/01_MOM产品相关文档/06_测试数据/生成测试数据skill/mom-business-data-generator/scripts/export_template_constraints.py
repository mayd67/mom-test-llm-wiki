from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

from generate_seed_workbooks import TEMPLATE_FILES, default_output_root_for_template_dir, find_template_dir

FILL_GUIDE_SHEET = "填写说明"
KEY_TEMPLATE_FILES = "模板文件"
KEY_GUIDE_NOTES = "填写说明"
KEY_ENUMS = "下拉枚举"
KEY_SPECIAL_RULES = "特殊规则"
KEY_ENUM_TIPS = "已验证的模板枚举注意事项"

SPECIAL_RULES = [
    "行政组织顶层公司的父组织编码建议使用 0。",
    "业务组织顶层公司的父组织编码建议使用 0。",
    "所有引用属性统一填写对象编码，不填写名称。",
    "凡是引用关系同时包含版本号时，版本号必须一并填写。",
    "当前 skill 主要用于 POC 验证时，默认采用最低安全口径：用户安全等级使用一般，数据密级使用公开。",
    "建议单个 Sheet 导入行数控制在 100 行以内，避免模板导入体验变差。",
]

ENUM_TIPS = [
    "行政组织类型只允许：公司、工厂、部门。",
    "业务组织-工厂类型只允许：机械加工专业、装配专业。",
    "工作中心-类型只允许：产线、设备组、人员组、外委、组织。",
    "工作中心-分类只允许：检验、加工。",
    "工艺路线-工艺专业只允许：机加、装配、热表、铸造、钣焊、锻造、通用。",
    "工艺路线工序/工序库-工序类型只允许：加工、检验、厂际转工、厂内转工、外委。",
    "工序专业类型只允许：机械加工专业、装配专业。",
    "当前模板中物料与工装工具的计量单位下拉仅提供：个。",
    "生产订单-订单类型只允许：标准。",
    "生产订单-业务状态只允许：初始、已展开、已释放、已开工、已完工。",
    "生产订单-控制状态只允许：正常、暂停、取消。",
    "备料清单-是否必须装入只允许：是、否。",
]


def normalize(value):
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    return str(value).strip()


def parse_validation_list(formula: str) -> list[str]:
    formula = normalize(formula)
    if formula.startswith('"') and formula.endswith('"'):
        return [normalize(item) for item in formula[1:-1].split(',') if normalize(item)]
    return []


def extract_constraints(template_dir: Path) -> dict[str, object]:
    data: dict[str, object] = {
        KEY_TEMPLATE_FILES: list(TEMPLATE_FILES),
        KEY_GUIDE_NOTES: {},
        KEY_ENUMS: {},
        KEY_SPECIAL_RULES: list(SPECIAL_RULES),
        KEY_ENUM_TIPS: list(ENUM_TIPS),
    }

    for workbook_name in TEMPLATE_FILES:
        workbook = load_workbook(template_dir / workbook_name)
        try:
            notes = workbook[FILL_GUIDE_SHEET]["A1"].value if FILL_GUIDE_SHEET in workbook.sheetnames else ""
            data[KEY_GUIDE_NOTES][workbook_name] = normalize(notes)
            workbook_constraints: dict[str, dict[str, list[str]]] = {}

            for worksheet in workbook.worksheets:
                headers = [normalize(cell.value) for cell in worksheet[1]]
                sheet_constraints: dict[str, list[str]] = {}
                if worksheet.data_validations:
                    for validation in worksheet.data_validations.dataValidation:
                        if validation.type != "list":
                            continue
                        allowed_values = parse_validation_list(validation.formula1)
                        if not allowed_values:
                            continue
                        for cell_range in validation.sqref.ranges:
                            if cell_range.min_col != cell_range.max_col:
                                continue
                            header = headers[cell_range.min_col - 1]
                            if not header:
                                continue
                            current = sheet_constraints.setdefault(header, [])
                            for item in allowed_values:
                                if item not in current:
                                    current.append(item)
                if sheet_constraints:
                    workbook_constraints[worksheet.title] = sheet_constraints

            data[KEY_ENUMS][workbook_name] = workbook_constraints
        finally:
            workbook.close()

    return data


def to_markdown(data: dict[str, object]) -> str:
    lines: list[str] = []
    lines.append("# MOM模板下拉枚举与填写约束字典")
    lines.append("")
    lines.append("## 用途")
    lines.append("")
    lines.append("本文件用于约束种子数据、Excel 填报数据和后续自动生成脚本，确保填写值命中模板真实下拉枚举，避免导入时报错。")
    lines.append("")
    lines.append("## 特殊规则")
    lines.append("")
    for item in data[KEY_SPECIAL_RULES]:
        lines.append(f"- {item}")
    lines.append("")
    lines.append("## 已验证的模板枚举注意事项")
    lines.append("")
    for item in data[KEY_ENUM_TIPS]:
        lines.append(f"- {item}")
    lines.append("")

    notes = data[KEY_GUIDE_NOTES]
    enums = data[KEY_ENUMS]
    for workbook_name in TEMPLATE_FILES:
        lines.append(f"## {workbook_name}")
        lines.append("")
        workbook_note = normalize(notes.get(workbook_name, ""))
        if workbook_note:
            lines.append("### 填写说明")
            lines.append("")
            for row in workbook_note.splitlines():
                row = normalize(row)
                if row:
                    lines.append(f"- {row}")
            lines.append("")

        workbook_enums = enums.get(workbook_name, {})
        for sheet_name, field_map in workbook_enums.items():
            lines.append(f"### Sheet：{sheet_name}")
            lines.append("")
            for field_name, values in field_map.items():
                joined = "、".join(values)
                lines.append(f"- `{field_name}`：{joined}")
            lines.append("")

    return "\n".join(lines).strip() + "\n"


def main() -> None:
    script_dir = Path(__file__).resolve().parent
    template_dir = find_template_dir(script_dir)
    if template_dir is None:
        raise FileNotFoundError("未找到 Excel 模板目录")

    output_dir = default_output_root_for_template_dir(template_dir) / "template_constraints"
    output_dir.mkdir(parents=True, exist_ok=True)

    data = extract_constraints(template_dir)
    markdown = to_markdown(data)

    json_path = output_dir / "template_constraint_dictionary.json"
    md_path = output_dir / "template_constraint_dictionary.md"
    ref_path = script_dir.parent / "references" / "template_constraint_dictionary.md"

    json_path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8-sig")
    md_path.write_text(markdown, encoding="utf-8-sig")
    ref_path.write_text(markdown, encoding="utf-8-sig")

    print(f"已生成 JSON：{json_path}")
    print(f"已生成 Markdown：{md_path}")
    print(f"已同步参考文档：{ref_path}")


if __name__ == "__main__":
    main()
